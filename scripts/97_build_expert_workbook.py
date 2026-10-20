# scripts/97_build_expert_workbook.py
"""生成发给专家的盲标工作簿。SEED=42。
S1：从审计后的 KB 中按 tier 各抽 40 条（共 120 条），盲标 (component, intent)。
S2：E1 候选池（e1_pool.csv），逐对判断是否相关，遗漏的对写在 missing_pairs 列。
S3：新 RAG 中有数值的 shall（A/B 级）25 条 + no-RAG 中 Grade X 的 shall 25 条，打乱顺序，隐去臂别和引用。
"""
from __future__ import annotations

import csv
import json
import random
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from cooling_rag_mg.evidence_taxonomy import COMPONENTS, INTENTS

KB = ROOT / "data/processed/evidence_units_v3_audited.jsonl"
E4 = ROOT / "data/search_results/exp_v4"
OUT = ROOT / "data/expert_round2"
README = [
    "本工作簿共四部分（S1–S4），请独立完成，不要查看项目代码或数据。预计 6–8 小时，可分多次完成。",
    "S1：阅读每条证据文本，从下拉框选出它最主要描述的部件(component)和工程方面(intent)，各选一个。",
    "S2：判断每个(部件, 方面)是否与该条原始需求直接相关：相关填1，不相关填0。若该需求还有其他相关对未列出，写在 missing_pairs，格式 comp:intent;comp:intent。",
    "S3：仅凭工程知识，判断每条 shall 的数值判据：correct(与规范或业界做法一致) / approx(±20%以内) / wrong(明显偏离) / unknown(超出把握)。",
    "S4：阅读每条证据文本（与抽取器所见相同，前1200字），对照右侧的抽取结果：n_correct 填抽取正确的数值个数，n_missed 填文本中明确出现但未被抽取的数值个数。",
    "部件与方面的含义见 S1 右侧的对照表。",
]


_NOTE_MARKERS = ("整理说明", "可支撑：", "可支撑:", "不能支撑")
_LABEL_NOTE = re.compile(r"针对\s*\S+\s*的\s*[^。]*?方面提供原文依据。?")


def annotator_text(t: str) -> str:
    """Original source wording only: drop curator notes, which name taxonomy labels."""
    t = t or ""
    cut = min([t.index(m) for m in _NOTE_MARKERS if m in t] or [len(t)])
    return _LABEL_NOTE.sub("", t[:cut]).strip()


def main() -> None:
    rng = random.Random(42)
    if (OUT / "expert_workbook.xlsx").exists() or (OUT / "expert_key.json").exists():
        raise SystemExit("Existing expert workbook/key protected; use an independently versioned preparation workflow")
    OUT.mkdir(parents=True, exist_ok=True)
    kb = [json.loads(l) for l in KB.open(encoding="utf-8")]
    kbi = {u["id"]: u for u in kb}
    wb = Workbook()
    ws = wb.active; ws.title = "README"
    for line in README:
        ws.append([line])
    key: dict = {"S1": {}, "S2": {}, "S3": {}}

    s1 = wb.create_sheet("S1_tagging")
    # no unit IDs in the sheet: IDs such as E-PUMP-RED-004 encode the label being annotated
    s1.append(["item", "text", "component", "intent", "", "component_ref", "desc", "intent_ref", "desc"])
    sample = []
    for tier in ("L1", "L2", "L3"):
        pool = sorted([u for u in kb if u["kb_tier"] == tier], key=lambda u: u["id"])
        sample += rng.sample(pool, min(40, len(pool)))
    rng.shuffle(sample)
    for n, u in enumerate(sample, 1):
        s1.append([n, annotator_text(u.get("text"))[:1500], "", ""])
        key["S1"][str(n)] = {"eu_id": u["id"], "component": u.get("primary_component"),
                             "intent": u.get("primary_intent"), "tier": u["kb_tier"]}
    for r, (c, d) in enumerate(COMPONENTS.items(), 2):
        s1.cell(r, 6, c); s1.cell(r, 7, d)
    for r, (i, d) in enumerate(INTENTS.items(), 2):
        s1.cell(r, 8, i); s1.cell(r, 9, d)
    dvc = DataValidation(type="list", formula1=f"=$F$2:$F${len(COMPONENTS)+1}")
    dvi = DataValidation(type="list", formula1=f"=$H$2:$H${len(INTENTS)+1}")
    s1.add_data_validation(dvc); s1.add_data_validation(dvi)
    dvc.add(f"C2:C{len(sample)+1}"); dvi.add(f"D2:D{len(sample)+1}")

    s2 = wb.create_sheet("S2_pairs")
    s2.append(["item", "or_id", "requirement_text", "component", "component_desc", "intent",
               "intent_desc", "relevant", "missing_pairs"])
    rows = list(csv.DictReader((E4 / "e1_pool.csv").open(encoding="utf-8-sig")))
    for n, r in enumerate(rows, 1):
        s2.append([n, r["or_id"], r["requirement_text"], r["component"], r["component_desc"],
                   r["intent"], r["intent_desc"], "", ""])
        key["S2"][str(n)] = [r["or_id"], r["component"], r["intent"]]
    dv01 = DataValidation(type="list", formula1='"1,0"'); s2.add_data_validation(dv01)
    dv01.add(f"H2:H{len(rows)+1}")

    num = re.compile(r"\d")
    def numeric(path, grades):
        d = json.loads(path.read_text(encoding="utf-8"))
        return [(rs["source_requirement_id"], e) for rs in d["result_sets"]
                for e in rs["engineering_requirements"]
                if e.get("_grade") in grades and num.search(str(e.get("measurable_criterion", "")))]
    rp, npool = numeric(E4 / "rag_v4_results.json", {"A", "B"}), numeric(E4 / "norag_v4_results.json", {"X"})
    m = min(25, len(rp), len(npool))  # 两臂等量；若 RAG 的 A+B 不足 25 条，按实际数量抽取，并在论文中写明 n
    rag, nor = rng.sample(rp, m), rng.sample(npool, m)
    items = [("RAG", s, e) for s, e in rag] + [("no-RAG", s, e) for s, e in nor]
    rng.shuffle(items)
    s3 = wb.create_sheet("S3_values")
    s3.append(["item", "shall_statement", "measurable_criterion", "verdict"])
    for n, (arm, sid, e) in enumerate(items, 1):
        s3.append([n, e.get("shall_statement", ""), str(e.get("measurable_criterion", "")), ""])
        key["S3"][str(n)] = {"arm": arm, "sid": sid, "grade": e.get("_grade")}
    dvv = DataValidation(type="list", formula1='"correct,approx,wrong,unknown"'); s3.add_data_validation(dvv)
    dvv.add(f"D2:D{len(items)+1}")

    s4 = wb.create_sheet("S4_extraction")
    s4.append(["item", "text_first_1200_chars", "extracted_values", "n_correct", "n_missed"])
    key["S4"] = {}
    rag_sets = json.loads((E4 / "rag_v4_results.json").read_text(encoding="utf-8"))["result_sets"]
    used = {}
    for rs in rag_sets:
        for eid, vals in (rs.get("_evidence_values") or {}).items():
            used.setdefault(eid, vals)
    by_tier = {t: sorted([i for i in used if i in kbi and kbi[i]["kb_tier"] == t]) for t in ("L1", "L2", "L3")}
    s4_items = []
    for t, q in (("L1", 17), ("L2", 17), ("L3", 16)):
        s4_items += [(t, i) for i in rng.sample(by_tier[t], min(q, len(by_tier[t])))]
    rng.shuffle(s4_items)
    for n, (t, i) in enumerate(s4_items, 1):
        vals = used[i]
        s4.append([n, (kbi[i].get("text") or "")[:1200], " | ".join(v.get("value_text", "") for v in vals), "", ""])
        key["S4"][str(n)] = {"eu_id": i, "tier": t, "n_extracted": len(vals)}

    wb.save(OUT / "expert_workbook.xlsx")
    (OUT / "expert_key.json").write_text(json.dumps(key, ensure_ascii=False, indent=1), encoding="utf-8")
    print("S1", len(sample), "S2", len(rows), "S3", len(items), "S4", len(s4_items), "->", OUT / "expert_workbook.xlsx")


if __name__ == "__main__":
    main()
