# -*- coding: utf-8 -*-
"""
Generate evidence1.docx — same format as evidence.docx —
containing all newly added evidence units (S60-S69).
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

EU_PATH = PROJECT_ROOT / "data" / "processed" / "evidence_units.jsonl"
OUT_PATH = PROJECT_ROOT / "data" / "raw" / "evidence1.docx"

# Source documents in the new range
NEW_SRC_PREFIXES = [f"S{n} " for n in range(60, 70)]


def _is_new(unit: dict) -> bool:
    doc = unit["metadata"].get("source_document", "")
    return any(doc.startswith(p) for p in NEW_SRC_PREFIXES)


def _list_field(val) -> str:
    if isinstance(val, list):
        return " / ".join(val) if val else ""
    return str(val) if val else ""


def _add_field(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    run_label = p.add_run(f"{label}：")
    run_label.bold = True
    p.add_run(value or "")


def _add_numbered_list(doc: Document, label: str, items) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    run = p.add_run(f"{label}：")
    run.bold = True
    if isinstance(items, list) and items:
        for i, item in enumerate(items, 1):
            if str(item).strip():
                doc.add_paragraph(f"{i}. {item}", style="Normal")
    elif isinstance(items, str) and items.strip():
        # might be newline-separated
        lines = [l.strip() for l in items.split("\n") if l.strip()]
        for i, line in enumerate(lines, 1):
            doc.add_paragraph(f"{i}. {line}", style="Normal")


def _build_supportable(unit: dict) -> list[str]:
    """Generate supportable requirements text from can_support list."""
    m = unit["metadata"]
    existing = m.get("supportable_requirements") or []
    if existing:
        return existing

    can = m.get("can_support") or []
    comps = m.get("component_objects") or []
    comp_str = "、".join(comps[:2]) if comps else "冷却水系统"

    support_map = {
        "capacity": f"{comp_str}的容量计算与设计余量要求。",
        "parameter_limit": f"{comp_str}相关温度、压力、流量参数的设定与确认。",
        "installation": f"{comp_str}安装规范与接口要求。",
        "test": f"{comp_str}安装后水压试验、功能试验和海试验收程序。",
        "commissioning_check": f"{comp_str}调试和试车验收项目。",
        "alarm": f"{comp_str}相关报警阈值的设定和功能验证要求。",
        "monitoring": f"{comp_str}运行参数的持续监控和仪表配置要求。",
        "operation": f"{comp_str}运行操作程序和注意事项。",
        "maintenance": f"{comp_str}定期维护保养要求。",
        "redundancy": f"{comp_str}冗余配置和备用切换要求。",
        "inspection": f"{comp_str}检验和见证要求。",
        "class_review": f"{comp_str}船级社审查和认证要求。",
        "calculation": f"{comp_str}相关水力或热力计算要求。",
        "safety": f"{comp_str}安全操作和危险防范要求。",
        "system_design": f"{comp_str}系统设计原则和配置要求。",
    }
    return [support_map.get(a, f"{comp_str}的{a}类需求。") for a in can if a in support_map]


def _build_unsupported(unit: dict) -> list[str]:
    """Generate unsupported requirements text."""
    m = unit["metadata"]
    existing = m.get("unsupported_requirements") or ""
    if existing and existing.strip():
        return [existing]

    app = m.get("applicability_level", "")
    src_type = m.get("source_type", "")

    limits = []
    if app == "class_rule":
        limits.append("具体实施细节和数值应以主机制造商技术文件和船级社批准图纸为最终依据。")
        limits.append("不同船级社在数值要求上可能存在细微差异，需按注册船级社规则版本确认。")
    elif app == "international_standard":
        limits.append("该标准提供框架要求，具体数值应以船级社规则和主机制造商技术文件为准。")
    elif app == "manufacturer_specific":
        limits.append("该参数值仅适用于该机型/配置，不可直接用于其他发动机。")
        limits.append("实际值应在调试阶段由厂家调试工程师最终确认。")
    else:
        limits.append("具体施工和数值要求应按项目技术规格书和船级社批准图纸确认。")

    return limits


def _add_entry(doc: Document, unit: dict) -> None:
    m = unit["metadata"]
    uid = unit["id"]
    topic = m.get("topic", "")

    # Heading 3
    h3 = doc.add_heading(f"{uid} {topic}", level=3)

    # Metadata fields
    _add_field(doc, "证据编号", uid)
    _add_field(doc, "来源文档", m.get("source_document", ""))
    _add_field(doc, "来源类型", m.get("source_type", ""))
    _add_field(doc, "来源链接", m.get("source_link", ""))
    _add_field(doc, "页码或位置", m.get("location", ""))
    _add_field(doc, "章节标题", m.get("section_title", ""))
    _add_field(doc, "系统层级", _list_field(m.get("system_scope")))
    _add_field(doc, "部件对象", _list_field(m.get("component_objects")))
    _add_field(doc, "相关提及对象", _list_field(m.get("related_objects")))
    _add_field(doc, "需求方面", _list_field(m.get("requirement_intents")))
    _add_field(doc, "工程阶段", _list_field(m.get("engineering_phase")))
    _add_field(doc, "适用工况", _list_field(m.get("applicable_condition")))
    _add_field(doc, "可支撑类型", _list_field(m.get("can_support")))
    _add_field(doc, "证据强度", m.get("evidence_strength", ""))
    _add_field(doc, "是否可泛化", m.get("is_generalizable", ""))
    _add_field(doc, "适用性等级", m.get("applicability_level", ""))
    _add_field(doc, "适用厂家", m.get("manufacturer", ""))
    _add_field(doc, "适用机型", m.get("engine_model", ""))
    _add_field(doc, "适用系统型式", _list_field(m.get("system_type")))
    _add_field(doc, "是否需要人工复核", m.get("need_human_review", ""))
    _add_field(doc, "适用对象", m.get("applicable_object", ""))
    _add_field(doc, "限制条件", m.get("limitation", ""))

    # 参数列表
    params = m.get("parameters") or []
    if params:
        p = doc.add_paragraph()
        p.style = doc.styles["Normal"]
        p.add_run("量化参数：").bold = True
        for param in params:
            name = param.get("name", "")
            val = param.get("value", "")
            unit_str = param.get("unit", "")
            cond = param.get("condition", "")
            line = f"{name} = {val} {unit_str}"
            if cond:
                line += f"  [{cond}]"
            doc.add_paragraph(f"  • {line}", style="Normal")

    # 原文
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.add_run("原文：").bold = True
    doc.add_paragraph(unit.get("text", ""), style="Normal")

    # 整理说明
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.add_run("整理说明：").bold = True
    doc.add_paragraph(m.get("normalized_summary", ""), style="Normal")

    # 可支撑的需求
    supportable = _build_supportable(unit)
    _add_numbered_list(doc, "可支撑的需求", supportable)

    # 不能支撑的需求
    unsupported = _build_unsupported(unit)
    _add_numbered_list(doc, "不能支撑的需求", unsupported)

    # 关键词
    kw_zh = m.get("keywords_zh") or []
    kw_en = m.get("keywords_en") or []
    all_kw = kw_zh + kw_en
    _add_field(doc, "建议关联关键词", "，".join(all_kw) if all_kw else
               "，".join((m.get("component_objects") or []) + (m.get("can_support") or [])[:3]))


def main():
    sys.stdout.reconfigure(encoding="utf-8")

    with open(EU_PATH, encoding="utf-8") as f:
        units = [json.loads(l) for l in f if l.strip()]

    new_units = [u for u in units if _is_new(u)]
    print(f"New units to include: {len(new_units)}")

    doc = Document()

    # Title
    title = doc.add_paragraph()
    title.style = doc.styles["Normal"]
    run = title.add_run("船舶主机冷却水系统 RAG 工程证据库 — 补充证据 V1")
    run.bold = True
    run.font.size = Pt(14)

    # Description
    doc.add_paragraph(
        "本文件为 evidence.docx 的补充版，收录针对证据库覆盖空白新增的权威来源证据单元，"
        "来源涵盖 IACS 统一要求（M3.1、M28、M51、M72、M73、E10）、"
        "DNV 船级社规则 Pt.4 Ch.3、中国船级社（CCS）第4篇第3章、"
        "以及 Wärtsilä 46DF 产品手册经验证参数。",
        style="Normal"
    )
    doc.add_paragraph(
        f"生成时间：2026-05-24；补充证据条数：{len(new_units)}。",
        style="Normal"
    )

    # Group by source document
    from collections import defaultdict
    by_source: dict[str, list] = defaultdict(list)
    for u in new_units:
        by_source[u["metadata"].get("source_document", "")].append(u)

    # Index
    doc.add_heading("证据来源目录", level=1)
    for i, (src, units_in_src) in enumerate(by_source.items(), 1):
        doc.add_paragraph(f"{i:02d}. {src}  ({len(units_in_src)} 条)", style="Normal")

    # Entries by source
    for src, units_in_src in by_source.items():
        doc.add_heading(src, level=1)
        for u in units_in_src:
            _add_entry(doc, u)

    doc.save(OUT_PATH)
    print(f"\nSaved: {OUT_PATH}")
    import os
    print(f"File size: {os.path.getsize(OUT_PATH) // 1024} KB")


if __name__ == "__main__":
    main()
