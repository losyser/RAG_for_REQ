from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document

from cooling_rag_mg.schema import normalize_v2_metadata, split_list
from cooling_rag_mg.text_utils import safe_text

# V2 语义化编号，例如 E-SYS-ACCESS-001、E-PUMP-CAP-001、E-SW-CLEAN-002。
# 同时兼容旧版 E00001。
ENTRY_RE = re.compile(r"^(?:#+\s*)?(E(?:-[A-Z0-9]+)+-\d{3}|E\d{5})(?:\s+(.+))?$")
KEY_VALUE_RE = re.compile(r"^([^：:]+)[：:](.*)$")

SECTION_KEYS = {
    "原文": "original_excerpt",
    "整理说明": "整理说明",
    "可支撑的需求": "可支撑的需求",
    "不能支撑的需求": "不能支撑的需求",
    "建议关联关键词": "建议关联关键词",
}

KNOWN_FIELD_KEYS = {
    "证据编号", "来源文档", "来源类型", "来源链接", "页码或位置", "位置", "章节标题",
    "系统层级", "部件对象", "相关提及对象", "需求方面", "工程阶段", "适用工况",
    "可支撑类型", "证据强度", "是否可泛化", "适用性等级", "适用厂家", "适用机型",
    "适用系统型式", "是否需要人工复核", "适用对象", "限制条件", "证据标题",
}


def read_docx_text(docx_path: str | Path) -> list[str]:
    doc = Document(str(docx_path))
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    # 有些 docx 可能使用表格存放证据，顺手兼容。
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append("：".join(cells) if len(cells) == 2 else " | ".join(cells))

    return lines


def parse_evidence_docx(docx_path: str | Path) -> list[dict[str, Any]]:
    lines = read_docx_text(docx_path)

    entries: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    active_section: str | None = None
    section_buffers: dict[str, list[str]] = {}

    def flush_current() -> None:
        nonlocal current, section_buffers
        if current is None:
            return
        for key, buf in section_buffers.items():
            current[key] = "\n".join([x for x in buf if safe_text(x)]).strip()
        entries.append(current)
        section_buffers = {}

    def start_section(key: str, value: str = "") -> None:
        nonlocal active_section, section_buffers
        active_section = key
        section_buffers.setdefault(key, [])
        if value:
            section_buffers[key].append(value)

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        m = ENTRY_RE.match(line)
        if m:
            flush_current()
            current = {
                "id": m.group(1).strip(),
                "topic": safe_text(m.group(2)),
            }
            active_section = None
            section_buffers = {}
            continue

        if current is None:
            continue

        # 处理“原文：xxx”这类段落标题。
        section_matched = False
        for label, key in SECTION_KEYS.items():
            if line == label or line in {f"{label}：", f"{label}:"}:
                start_section(key)
                section_matched = True
                break
            if line.startswith(f"{label}：") or line.startswith(f"{label}:"):
                value = line.split("：", 1)[1] if "：" in line else line.split(":", 1)[1]
                start_section(key, value.strip())
                section_matched = True
                break
        if section_matched:
            continue

        kv = KEY_VALUE_RE.match(line)
        if kv:
            key = kv.group(1).strip()
            value = kv.group(2).strip()
            if key in SECTION_KEYS:
                start_section(SECTION_KEYS[key], value)
                continue
            if key in KNOWN_FIELD_KEYS:
                # 遇到普通字段后，结束多行段落收集。
                active_section = None
                current[key] = value
                continue
            # active_section 内的英文冒号、邮箱、比例等不应误判为字段。
            if active_section:
                section_buffers.setdefault(active_section, []).append(line)
                continue
            current[key] = value
            continue

        if active_section:
            section_buffers.setdefault(active_section, []).append(line)

    flush_current()

    return [normalize_raw_entry(e) for e in entries]


def normalize_raw_entry(e: dict[str, Any]) -> dict[str, Any]:
    evidence_id = safe_text(e.get("证据编号")) or safe_text(e.get("id"))
    topic = safe_text(e.get("证据标题")) or safe_text(e.get("topic"))

    metadata = {
        "topic": topic,
        "source_document": safe_text(e.get("来源文档")),
        "source_type": safe_text(e.get("来源类型")),
        "source_link": safe_text(e.get("来源链接")),
        "location": safe_text(e.get("页码或位置") or e.get("位置")),
        "section_title": safe_text(e.get("章节标题")),
        "system_scope": split_list(e.get("系统层级")),
        "component_objects": split_list(e.get("部件对象")),
        "related_objects": split_list(e.get("相关提及对象")),
        "requirement_intents": split_list(e.get("需求方面")),
        "engineering_phase": split_list(e.get("工程阶段")),
        "applicable_condition": split_list(e.get("适用工况")),
        "can_support": split_list(e.get("可支撑类型")),
        "evidence_strength": safe_text(e.get("证据强度")) or "unknown",
        "is_generalizable": safe_text(e.get("是否可泛化")) or "unknown",
        "applicability_level": safe_text(e.get("适用性等级")) or "unknown",
        "manufacturer": safe_text(e.get("适用厂家")),
        "engine_model": safe_text(e.get("适用机型")),
        "system_type": split_list(e.get("适用系统型式")),
        "need_human_review": safe_text(e.get("是否需要人工复核")) or "no",
        "applicable_object": safe_text(e.get("适用对象")),
        "limitation": safe_text(e.get("限制条件")),
        "normalized_summary": safe_text(e.get("整理说明")),
        "supportable_requirements": split_numbered_block(e.get("可支撑的需求")),
        "unsupported_requirements": safe_text(e.get("不能支撑的需求")),
        "keywords_zh": split_list(e.get("建议关联关键词")),
        "keywords_en": [],
        "parameters_raw": safe_text(e.get("参数")),
        "parameters": [],
        "parameter_name": [],
        "value": [],
        "unit": [],
        "condition": [],
        "is_noise": False,
        "noise_reason": "",
    }
    metadata = normalize_v2_metadata(metadata)

    return {
        "id": evidence_id,
        "topic": topic,
        "original_excerpt": safe_text(e.get("original_excerpt")),
        "metadata": metadata,
    }


def split_numbered_block(value: Any) -> list[str]:
    text = safe_text(value)
    if not text:
        return []
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    cleaned: list[str] = []
    for line in lines:
        line = re.sub(r"^\d+[）\).、]\s*", "", line).strip()
        if line:
            cleaned.append(line)
    return cleaned
