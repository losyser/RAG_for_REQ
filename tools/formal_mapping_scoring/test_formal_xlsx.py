"""Read-only originals; all generated labels confined to disposable synthetic copies."""
import hashlib,json,subprocess,sys,tempfile,unittest
from pathlib import Path
from openpyxl import load_workbook
import excel_pilot as reader
import score_formal_review as scorer
ROOT=Path(r'D:\cooling_rag_multigranular\cooling_rag_multigranular')
DESKTOP=Path(r'C:\Users\yando\Desktop\需求映射正式审核_8源')
PACKAGE=DESKTOP/'formal_8_sources.assignment.json'
MANIFEST=ROOT/'revision_artifacts/expert_mapping_formal_20261017/coordinator/FROZEN_SAMPLE.json'
PRED=ROOT/'data/search_results/exp_v4/e1_predictions.json'

class FormalExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.out=Path(self.tmp.name)
        self.package=scorer.strict_json(PACKAGE)
        self.template=DESKTOP/'需求映射正式审核_8源_专家A.xlsx'
    def tearDown(self):self.tmp.cleanup()
    def edited_copy(self,edit):
        book=load_workbook(self.template);edit(book)
        target=self.out/'SYNTHETIC_ONLY.xlsx';book.save(target);book.close();return target
    def complete_copy(self):
        def fill(book):
            home=book.worksheets[0];home['B6']='SYNTHETIC_ONLY_NOT_EXPERT'
            for i,s in enumerate(self.package['sources']):
                start=17+14*i
                for offset,value in [(2,'是'),(3,'首判完成后查看'),(4,'2026-10-17T10:01:00Z'),(5,'是'),(6,'是'),(7,10),(8,0),(11,'2026-10-17T10:00:00Z')]:home.cell(start+offset,2,value)
                outside=next((c,intent) for c,intents in self.package['taxonomy'].items() for intent in intents if c+':'+intent not in {p['id'] for p in s['candidates']})
                sheet=book.worksheets[3]
                for col,value in [(1,s['id']),(2,':'.join(outside)),(5,'相关'),(6,'SYNTHETIC TEST ONLY')]:sheet.cell(9+i,col,value)
            for r in range(9,84):book.worksheets[2].cell(r,5,'相关')
        return self.edited_copy(fill)
    def cli(self,record):
        output=self.out/('result_'+str(len(list(self.out.glob('result*'))))+'.json')
        run=subprocess.run([sys.executable,'-I',str(Path(scorer.__file__)),'--package',str(PACKAGE),'--manifest',str(MANIFEST),'--predictions',str(PRED),'--records',str(record),'--expected-reviewers','1','--output',str(output)],capture_output=True,text=True)
        self.assertEqual(run.returncode,0,run.stderr);return scorer.strict_json(output)
    def test_both_originals_blank_same_sample_and_hash(self):
        versions=[]
        for slot in ('A','B'):
            path=DESKTOP/f'需求映射正式审核_8源_专家{slot}.xlsx';before=hashlib.sha256(path.read_bytes()).hexdigest()
            r=reader.read_workbook(path,PACKAGE);versions.append(r['package_version'])
            self.assertEqual(r['reviewer_slot'],slot);self.assertEqual(len(r['candidate_rows']),75)
            self.assertTrue(all(x['raw_verdict'] is None for x in r['candidate_rows']))
            self.assertFalse(r['pending_pair_rows']);self.assertTrue(all(not x for x in r['pair_rows'].values()))
            for raw in r['raw_source_fields'].values():self.assertTrue(all(v is None for k,v in raw.items() if k!='original_text'))
            self.assertEqual(list(r['web_answer_fields']),scorer.strict_json(MANIFEST)['selected_source_ids'])
            self.assertEqual(before,hashlib.sha256(path.read_bytes()).hexdigest())
        self.assertEqual(versions,[self.package['version']]*2)
    def test_actual_cli_blank_has_no_f1(self):
        r=self.cli(self.template);self.assertEqual(r['status'],'blocked');self.assertIsNone(r['scores'])
        self.assertEqual(r['coverage'][0]['candidate_denominator'],75);self.assertEqual(r['coverage'][0]['definitive_candidate_count'],0)
    def test_actual_cli_synthetic_complete_pool_outside_recall(self):
        r=self.cli(self.complete_copy());self.assertEqual(r['status'],'scored',r['reasons'])
        predictions=scorer.strict_json(PRED)
        for method in predictions:
            count=sum(len(predictions[method][sid]) for sid in r['source_ids'])
            self.assertEqual(r['scores'][method]['micro']['tp'],count)
            self.assertEqual(r['scores'][method]['micro']['fp'],0)
            self.assertEqual(r['scores'][method]['micro']['fn'],83-count)
            self.assertAlmostEqual(r['scores'][method]['micro']['recall'],count/83)
    def test_duplicate_excel_candidate_ids_rejected(self):
        def duplicate(b):
            for col in (1,2):b.worksheets[2].cell(10,col,b.worksheets[2].cell(9,col).value)
        with self.assertRaises(ValueError):reader.read_workbook(self.edited_copy(duplicate),PACKAGE)
    def test_excel_wrong_version_rejected(self):
        with self.assertRaises(ValueError):reader.read_workbook(self.edited_copy(lambda b:setattr(b.worksheets[0]['B133'],'value','wrong')),PACKAGE)
    def test_duplicate_readback_rows_rejected(self):
        r=reader.read_workbook(self.template,PACKAGE);r['candidate_rows'][1]=r['candidate_rows'][0].copy()
        with self.assertRaises(ValueError):scorer.normalize_record(r,self.package)

if __name__=='__main__':unittest.main(verbosity=2)
