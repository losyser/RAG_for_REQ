"""Read an approved pilot or formal XLSX into a separate, loss-preserving JSON record. Never writes web tasks."""
import argparse
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import uuid
from zipfile import ZipFile
from openpyxl import load_workbook

spec=importlib.util.spec_from_file_location('review_server',Path(__file__).with_name('server.py'))
review=importlib.util.module_from_spec(spec);spec.loader.exec_module(review)
SHEETS=['0规则与计时','1独立首判','2候选判断','3遗漏补充','4完整分类目录']
STATES={None:'unanswered','':'unanswered','未答':'unanswered','相关':'relevant','不相关':'irrelevant','不确定':'uncertain'}
EXPOSURE={None,'','尚未查看','首判完成后查看','首判前已查看','不确定'}
HEADERS=['源ID','稳定配对ID','部件中文（自动）','工程方面（自动）','判断','中文理由／疑问']
HOME_FIELDS=['original_text','first_pass_done','candidate_exposure','first_candidate_view_time','coverage_reviewed','declared_complete','total_minutes','interruption_minutes','questions','source_note','first_pass_completed_time']

def require(condition,message):
    if not condition:raise ValueError(message)

def value(cell):
    require(cell.data_type!='f',f'{cell.parent.title}!{cell.coordinate} 为公式；输入栏仅接受实际值，请保留原意见为文本。')
    return cell.value

def string(raw,label,limit=20000):
    require(raw is None or isinstance(raw,str),f'{label} 应为文本')
    require(raw is None or len(raw)<=limit,f'{label} 超过长度限制')
    return raw or ''

def read_workbook(workbook_path,package_path):
    require(workbook_path.stat().st_size<=8_000_000,'工作簿过大，请提供本试做模板的.xlsx文件')
    with ZipFile(workbook_path) as archive:
        require(not any(n.startswith('xl/externalLinks/') or n.endswith('vbaProject.bin') for n in archive.namelist()),'不支持外部链接或宏')
    package=json.loads(package_path.read_text(encoding='utf8'))
    review.validate_package(package)
    formal=package['kind']=='formal'
    require(package['kind'] in ('pilot','formal'),'仅支持已确认评审包')
    meta_start=19+14*len(package['sources'])
    book=load_workbook(workbook_path,data_only=False,keep_links=False)
    try:
        require(book.sheetnames==SHEETS,'工作表名称或数量改变，请保留模板结构')
        home=book[SHEETS[0]]
        slot=value(home.cell(meta_start+8,2)) if formal else None
        require(not formal or slot in ('A','B'),'独立专家槽位无效')
        metadata=['mapping-review-excel-formal-v1' if formal else 'mapping-review-excel-pilot-v1',package['kind']+'-'+package['version'][:16]+('-'+slot if formal else ''),package['version'],package['protocol'],package['provenance']['source_sha256'],package['provenance']['taxonomy_sha256'],package['kind'],package_path.name]
        require([value(home.cell(r,2)) for r in range(meta_start,meta_start+8)]==metadata,'来源版本或pilot元数据改变，拒绝混用')
        reviewer=string(value(home['B6']),'专家代号',80)
        legal={c+':'+i:(c,i,package['components'][c],package['intents'][i]) for c,items in package['taxonomy'].items() for i in items}
        tax=book[SHEETS[4]]
        require(tax.max_row==187,'完整目录行数改变')
        seen=set()
        for row in tax.iter_rows(min_row=5,max_row=187,max_col=5):
            pid,cn,inn,c,i=[value(cell) for cell in row]
            require(pid in legal and pid not in seen and (c,i,cn,inn)==legal[pid],'完整目录ID、说明或合法配对改变')
            seen.add(pid)
        require(seen==set(legal),'完整目录不完整')
        answers=review.blank_answers(package)
        raw_sources={};issues=[];candidate_rows=[];pair_rows={SHEETS[1]:[],SHEETS[3]:[]};pending=[]
        if not reviewer.strip():issues.append('专家代号未填写')
        sources={s['id']:s for s in package['sources']}
        for index,source in enumerate(package['sources']):
            sid=source['id'];start=17+index*14
            require(value(home.cell(start,1))==sid,'首页源ID改变')
            raw={k:value(home.cell(start+1+i,2)) for i,k in enumerate(HOME_FIELDS)}
            require(raw['original_text']==source['text'],'原始需求中文被修改')
            for sheet_name in ([] if formal else SHEETS[1:4]):
                ws=book[sheet_name];require(value(ws.cell(4+index,1))==sid and value(ws.cell(4+index,2))==source['text'],sheet_name+' 的源ID或原文改变')
            for key in ('first_pass_done','coverage_reviewed','declared_complete'):
                require(raw[key] in (None,'','是','否'),sid+' '+key+' 只能是/否/空白')
            require(raw['candidate_exposure'] in EXPOSURE,sid+' 候选暴露状态无效')
            for key in ('first_candidate_view_time','first_pass_completed_time','questions','source_note'):
                string(raw[key],sid+' '+key,4000)
            for key in ('total_minutes','interruption_minutes'):
                x=raw[key];require(x is None or (type(x) in (int,float) and math.isfinite(x) and x>=0),sid+' '+key+' 必须为非负数或空白')
            if raw['total_minutes'] is not None and raw['interruption_minutes'] is not None:
                require(raw['interruption_minutes']<=raw['total_minutes'],sid+' 中断时长不能大于总耗时')
            answers[sid]['first_pass_done']=raw['first_pass_done']=='是'
            answers[sid]['coverage_reviewed']=raw['coverage_reviewed']=='是'
            answers[sid]['note']=raw['source_note'] or ''
            raw_sources[sid]=raw
            if formal:
                answers[sid]['review_details']={'exposure':{None:'uncertain','':'uncertain','尚未查看':'not_viewed','首判完成后查看':'after_first_pass','首判前已查看':'before_first_pass','不确定':'uncertain'}[raw['candidate_exposure']],
                    'first_pass_at':raw['first_pass_completed_time'] or '', 'first_candidate_at':raw['first_candidate_view_time'] or '',
                    'total_minutes':raw['total_minutes'],'interruption_minutes':raw['interruption_minutes'],
                    'complete':{'是':'yes','否':'no',None:'unanswered','':'unanswered'}[raw['declared_complete']],'questions':raw['questions'] or ''}
            for key in ('first_pass_done','coverage_reviewed','declared_complete'):
                if raw[key]!='是':issues.append(sid+' 尚未确认 '+key)
            if raw['total_minutes'] is None:issues.append(sid+' 未记录总耗时')
            if raw['candidate_exposure']!='首判完成后查看':issues.append(sid+' 独立首判流程未确认，候选暴露状态='+str(raw['candidate_exposure']))
        # Stable source/pair IDs, not row position, bind every candidate response.
        ws=book[SHEETS[2]]
        require([value(ws.cell(8,c)) for c in range(1,7)]==['源ID','稳定配对ID','部件中文','工程方面中文','判断','中文理由／疑问'],'候选表头改变')
        require(ws.max_row<=2000 and ws.max_column<=6,'候选页含模板范围外的内容')
        seen=set()
        for row in ws.iter_rows(min_row=9,max_row=ws.max_row,max_col=6):
            sid,pid,cn,inn,verdict,note=[value(c) for c in row]
            if all(x is None for x in (sid,pid,cn,inn,verdict,note)):continue
            require(sid in sources and pid in answers[sid]['candidates'],'未知候选源ID或配对ID')
            require((sid,pid) not in seen,'重复候选配对')
            require((cn,inn)==legal[pid][2:],'候选中文说明改变')
            require(verdict in STATES,'候选判断值无效')
            string(note,'候选中文理由',4000)
            seen.add((sid,pid));answers[sid]['candidates'][pid]=STATES[verdict]
            candidate_rows.append({'source_id':sid,'pair_id':pid,'raw_verdict':verdict,'verdict':STATES[verdict],'note':note})
        require(seen=={(sid,pid) for sid,a in answers.items() for pid in a['candidates']},'候选行缺失')
        for sheet_name,field in [(SHEETS[1],'initial_pairs'),(SHEETS[3],'missing_pairs')]:
            ws=book[sheet_name]
            require(ws.max_row<=2000 and ws.max_column<=6,'配对页超出支持的2000行/6列范围')
            require([value(ws.cell(8,c)) for c in range(1,7)]==HEADERS,'配对表头改变')
            seen=set()
            for r in range(9,ws.max_row+1):
                sid,pid,verdict,note=[value(ws.cell(r,c)) for c in (1,2,5,6)]
                if all(x is None for x in (sid,pid,verdict,note)):continue
                require(sid in sources or sid in (None,''),'未知源ID')
                require(pid in legal or pid in (None,''),'非法配对ID')
                require(verdict in (None,'','相关','不确定','未答'),'首判/补漏判断只能相关、不确定或未答')
                string(note,'首判/补漏中文意见',1000)
                raw={'sheet_row':r,'source_id':sid,'pair_id':pid,'raw_verdict':verdict,'verdict':STATES[verdict],'note':note}
                pair_rows[sheet_name].append(raw)
                if not sid or not pid or STATES[verdict]=='unanswered':
                    pending.append({'sheet':sheet_name,**raw});issues.append(f'{sheet_name} 第{r}行尚未填写完整，原值已保留');continue
                require((sid,pid) not in seen,'首判/补漏页有重复配对')
                require(field=='initial_pairs' or pid not in answers[sid]['candidates'],'补漏配对已在候选中，请在候选原行记录')
                seen.add((sid,pid));c,i,_,_=legal[pid]
                answers[sid][field].append({'component':c,'intent':i,'verdict':STATES[verdict],'note':note or ''})
        for sid,a in answers.items():
            if 'unanswered' in a['candidates'].values():issues.append(sid+' 存在未答候选')
            answered=any(v!='unanswered' for v in a['candidates'].values())
            if answered and raw_sources[sid]['candidate_exposure'] in (None,'','尚未查看'):
                issues.append(sid+' 已有候选判断但暴露状态未更新')
            if raw_sources[sid]['first_pass_done']=='是' and not raw_sources[sid]['first_pass_completed_time']:
                issues.append(sid+' 首判完成时间未填写')
            if answered and not raw_sources[sid]['first_candidate_view_time']:
                issues.append(sid+' 首次查看候选时间未填写')
        web_valid=True
        try:review.validate_answers(package,answers)
        except review.ReviewError as error:web_valid=False;issues.append('网页字段结构校验未通过：'+str(error))
        return {
            'format':'mapping-review-excel-import-v1','record_id':uuid.uuid4().hex,'kind':package['kind'],'excluded_from_formal_scoring':not formal,'reviewer_slot':slot,
            'collection_mode':'excel','stage_order':'both-source-first-pass-before-candidate-sheet',
            'reviewer':reviewer,'package_version':package['version'],'protocol':package['protocol'],
            'provenance':{'workbook_name':workbook_path.name,'workbook_sha256':hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
                          'package_sha256':hashlib.sha256(package_path.read_bytes()).hexdigest(),'source_sha256':package['provenance']['source_sha256'],'taxonomy_sha256':package['provenance']['taxonomy_sha256']},
            'raw_source_fields':raw_sources,'candidate_rows':candidate_rows,'pair_rows':pair_rows,'pending_pair_rows':pending,
            'web_answer_fields':answers,'web_answer_schema_valid':web_valid,
            'issues':issues,'annotation_complete':not issues,'pilot_review_complete':not formal and not issues,'gold_ready':False,
            'limitations':['Excel阶段和候选暴露为专家自报，不能证明强盲法或首判未被回改。','独立读回记录，不是网页备份，未写入或覆盖任何网页任务。','候选逐行意见、计时和暴露状态保存在本记录扩展字段中；网页旧答案结构没有这些字段，禁止有损覆盖。']}
    finally:book.close()

def convert(workbook_path,package_path,output):
    require(not output.exists(),'输出已存在，拒绝覆盖')
    record=read_workbook(workbook_path,package_path)
    with output.open('x',encoding='utf8') as handle:json.dump(record,handle,ensure_ascii=False,indent=2)
    return record

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook',type=Path);parser.add_argument('--package',type=Path,required=True);parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    try:
        result=convert(args.workbook,args.package,args.output)
        print(json.dumps({'output':str(args.output),'issues':len(result['issues']),'annotation_complete':result['annotation_complete'],'gold_ready':False},ensure_ascii=True))
    except (ValueError,OSError) as error:parser.exit(2,str(error)+'\n')
