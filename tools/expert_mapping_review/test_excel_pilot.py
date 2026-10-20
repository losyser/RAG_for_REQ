"""Synthetic edits in temporary copies only. No expert opinions are generated."""
import copy
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook

spec=importlib.util.spec_from_file_location('excel_converter',Path(__file__).with_name('excel_pilot.py'))
converter=importlib.util.module_from_spec(spec);spec.loader.exec_module(converter)
TEMPLATE=Path(os.environ['PILOT_XLSX_TEMPLATE'])
PACKAGE=Path(os.environ['PILOT_JSON_PACKAGE'])

class ExcelPilotTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.folder=Path(self.tmp.name)
        self.package=json.loads(PACKAGE.read_text(encoding='utf8'))
    def tearDown(self):self.tmp.cleanup()
    def change(self,callback):
        book=load_workbook(TEMPLATE);callback(book)
        path=self.folder/'synthetic_test_only.xlsx';book.save(path);book.close();return path
    def read(self,path):return converter.read_workbook(path,PACKAGE)
    def test_blank_template_and_native_features(self):
        result=self.read(TEMPLATE)
        self.assertEqual(result['kind'],'pilot');self.assertFalse(result['gold_ready']);self.assertFalse(result['pilot_review_complete'])
        self.assertTrue(all(r['raw_verdict'] is None and r['verdict']=='unanswered' for r in result['candidate_rows']))
        b=load_workbook(TEMPLATE)
        self.assertEqual(b.sheetnames,converter.SHEETS)
        self.assertEqual(b['4完整分类目录'].max_row,187)
        self.assertEqual(b['2候选判断'].freeze_panes,'C9')
        self.assertTrue(b['2候选判断']['D9'].alignment.wrap_text)
        self.assertGreater(b['2候选判断'].column_dimensions['D'].width,40)
        self.assertTrue(any('不确定' in (v.formula1 or '') and '未答' in (v.formula1 or '') for v in b['2候选判断'].data_validations.dataValidation))
        self.assertIn('ValidPairs',b.defined_names)
        self.assertTrue(all(s.sheet_state=='visible' for s in b))
        b.close()
    def test_roundtrip_values_uncertainty_notes_timing_and_missing_pairs(self):
        s=self.package['sources'][0];candidates={p['id'] for p in s['candidates']}
        legal=[c+':'+i for c,intents in self.package['taxonomy'].items() for i in intents]
        extra=[p for p in legal if p not in candidates]
        def fill(b):
            h=b['0规则与计时'];h['B6']='隔离测试，非专家';h['B19']='是';h['B20']='首判完成后查看';h['B21']='2026-10-17 14:10';h['B28']='2026-10-17 14:00';h['B24']=12.5;h['B25']=0;h['B26']='测试疑问：不能据此评分';h['B27']='原中文意见\n第二行保留'
            c=b['2候选判断']
            for row,v in [(9,'相关'),(10,'不相关'),(11,'不确定'),(12,'未答')]:c.cell(row,5,v)
            c['F11']='人工构造测试：不确定原因'
            for name,pid,v in [('1独立首判',extra[0],'不确定'),('3遗漏补充',extra[1],'相关')]:
                w=b[name];w['A9']='OR002';w['B9']=pid;w['E9']=v;w['F9']='测试补充意见\n保留换行'
        result=self.read(self.change(fill))
        a=result['web_answer_fields']['OR002']
        self.assertEqual(list(a['candidates'].values())[:5],['relevant','irrelevant','uncertain','unanswered','unanswered'])
        self.assertEqual(a['initial_pairs'][0]['verdict'],'uncertain');self.assertEqual(a['missing_pairs'][0]['verdict'],'relevant')
        self.assertEqual(a['note'],'原中文意见\n第二行保留')
        self.assertEqual(result['candidate_rows'][2]['note'],'人工构造测试：不确定原因')
        self.assertEqual(result['raw_source_fields']['OR002']['total_minutes'],12.5)
        self.assertEqual(result['raw_source_fields']['OR002']['interruption_minutes'],0)
        self.assertTrue(result['web_answer_schema_valid']);self.assertFalse(result['gold_ready'])
    def test_pending_pair_row_preserved(self):
        def fill(b):b['1独立首判']['A9']='OR002';b['1独立首判']['F9']='尚未决定配对，保留原意见'
        r=self.read(self.change(fill));self.assertEqual(r['pending_pair_rows'][0]['note'],'尚未决定配对，保留原意见');self.assertFalse(r['pilot_review_complete'])
    def test_invalid_inputs_and_source_tampering_rejected(self):
        for sheet,cell,bad in [('0规则与计时','B49','wrong-version'),('0规则与计时','B18','改写原文'),('2候选判断','E9','看起来不错'),('2候选判断','B9','pump:fake'),('4完整分类目录','C5','修改含义'),('0规则与计时','B24',-2),('2候选判断','F9','=1+1')]:
            with self.subTest(cell=cell,bad=bad):
                def edit(b):b[sheet][cell]=bad
                with self.assertRaises(ValueError):self.read(self.change(edit))
    def test_duplicate_missing_and_candidate_rows_rejected(self):
        def dup(b):
            s=b['2候选判断']
            for col in range(1,7):s.cell(10,col,s.cell(9,col).value)
        with self.assertRaises(ValueError):self.read(self.change(dup))
        def missing(b):
            s=b['2候选判断'];m=b['3遗漏补充'];m['A9']=s['A9'].value;m['B9']=s['B9'].value;m['E9']='相关'
        with self.assertRaises(ValueError):self.read(self.change(missing))
    def test_candidate_sort_uses_stable_ids(self):
        def reorder(b):
            s=b['2候选判断'];rows=[list(row) for row in s.iter_rows(min_row=9,max_row=34,values_only=True)]
            rows[0][4]='不确定'
            for rownum,row in enumerate(reversed(rows),9):
                for col,v in enumerate(row,1):s.cell(rownum,col).value=v
        r=self.read(self.change(reorder));p=self.package['sources'][0]['candidates'][0]['id']
        self.assertEqual(r['web_answer_fields']['OR002']['candidates'][p],'uncertain')
    def test_stage_exposure_and_incomplete_checks(self):
        def fill(b):
            b['2候选判断']['E9']='相关';b['0规则与计时']['B20']='首判前已查看';b['0规则与计时']['B23']='是'
        r=self.read(self.change(fill));self.assertFalse(r['web_answer_schema_valid']);self.assertFalse(r['pilot_review_complete']);self.assertTrue(any('候选暴露' in x for x in r['issues']))
    def test_completed_pilot_never_gold_and_no_overwrite(self):
        def fill(b):
            h=b['0规则与计时'];h['B6']='人工构造完整测试'
            for start in (17,31):
                for offset in (2,5,6):h.cell(start+offset,2,'是')
                h.cell(start+3,2,'首判完成后查看');h.cell(start+4,2,'2026-10-17 15:00');h.cell(start+11,2,'2026-10-17 14:00');h.cell(start+7,2,20);h.cell(start+8,2,0)
            for row in range(9,35):b['2候选判断'].cell(row,5,'不相关')
        r=self.read(self.change(fill));self.assertTrue(r['pilot_review_complete']);self.assertFalse(r['gold_ready'])
        out=self.folder/'record.json';converter.convert(TEMPLATE,PACKAGE,out);before=out.read_bytes()
        with self.assertRaises(ValueError):converter.convert(TEMPLATE,PACKAGE,out)
        self.assertEqual(out.read_bytes(),before)

if __name__=='__main__':unittest.main()
