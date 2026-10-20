import copy
import pytest
from openpyxl import Workbook
from cooling_rag_mg.expert_integrity import make_contract,validate_workbook,verify_sources,SHEETS

def fixture():
    wb=Workbook(); wb.remove(wb.active)
    headers={'S1':['item','text','component','intent'],'S2':['item','or_id','requirement_text','component','intent','relevant','missing_pairs'],'S3':['item','shall_statement','measurable_criterion','verdict'],'S4':['item','text_first_1200_chars','extracted_values','n_correct','n_missed']}
    data={'S1':[1,'source',None,None],'S2':[1,'OR1','requirement','pump','cooling',None,None],'S3':[1,'shall','80°C',None],'S4':[1,'source','80°C',None,None]}
    for s,n in SHEETS.items(): ws=wb.create_sheet(n);ws.append(headers[s]);ws.append(data[s])
    key={s:{'1':{}} for s in SHEETS};key['S2']['1']=['OR1','pump','cooling']
    return wb,make_contract(wb,key,{})

@pytest.mark.parametrize('section',['S1','S2','S3','S4'])
def test_immutable_text_tamper_blocked(section):
    wb,c=fixture();wb[SHEETS[section]].cell(2,2,'altered')
    assert validate_workbook(wb,c)

@pytest.mark.parametrize('mutation',['delete','duplicate','blank_id','extra','header','formula'])
def test_sample_structure_blocked(mutation):
    wb,c=fixture();ws=wb['S3_values']
    if mutation=='delete':ws.delete_rows(2)
    elif mutation=='duplicate':ws.append([1,'shall','80°C',None])
    elif mutation=='blank_id':ws['A2']=None
    elif mutation=='extra':ws.append([2,'shall','80°C',None])
    elif mutation=='header':ws['B1']='new header'
    else:ws['D2']='=1'
    assert validate_workbook(wb,c)

def test_blanks_integrity_only_does_not_authorize_scoring():
    wb,c=fixture();assert validate_workbook(wb,c)==[]
    # Annotation completeness is separately blocked by validate_expert_rows.
    from cooling_rag_mg.revision_audit import validate_expert_rows
    assert validate_expert_rows({'S3':[{'item':1,'verdict':None}]},{'S3':{'1':{}}})

def test_key_mismatch_and_manifest_tamper_blocked(tmp_path):
    wb,c=fixture(); key={s:{'1':{}} for s in SHEETS};key['S2']['1']=['OR2','pump','cooling']
    with pytest.raises(ValueError):make_contract(wb,key,{})
    assert verify_sources(tmp_path,c,{'files':[]})

def test_annotation_cells_can_change_without_altering_source():
    wb,c=fixture();wb['S3_values']['D2']='unknown'
    assert validate_workbook(wb,c)==[]
