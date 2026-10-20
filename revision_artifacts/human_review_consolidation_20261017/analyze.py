"""Post-hoc descriptive analysis, not a replacement for the frozen complete-gold scorer.
Reads original XLSX bytes, never writes workbooks or changes human labels.
"""
from pathlib import Path
from collections import Counter
import json,csv,hashlib,importlib.util,shutil,itertools
from openpyxl import load_workbook
B=Path(__file__).parent
ROOT=Path(r'D:\cooling_rag_multigranular\cooling_rag_multigranular')
PKG=(B/'inputs/formal_8_sources.assignment.json') if (B/'inputs/formal_8_sources.assignment.json').exists() else ROOT/'revision_artifacts/expert_mapping_formal_20261017/expert_packet/formal_8_sources.assignment.json'
FROZEN=(B/'inputs/FROZEN_SAMPLE.json') if (B/'inputs/FROZEN_SAMPLE.json').exists() else ROOT/'revision_artifacts/expert_mapping_formal_20261017/coordinator/FROZEN_SAMPLE.json'
PRED=(B/'inputs/e1_predictions.json') if (B/'inputs/e1_predictions.json').exists() else ROOT/'data/search_results/exp_v4/e1_predictions.json'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def save(path,data):path.write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str),encoding='utf8')
def metric(pred,positive,mask):
    pp=pred&mask;g=positive&mask
    tp=len(pp&g);fp=len(pp-g);fn=len(g-pp);tn=len(mask-(pp|g))
    return {'tp':tp,'fp':fp,'fn':fn,'tn':tn,'precision':tp/(tp+fp) if tp+fp else None,'recall':tp/(tp+fn) if tp+fn else None,'f1':2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else None,'mask_n':len(mask),'positive_n':len(g),'predictions_evaluated':len(pp),'predictions_outside_mask':len(pred-mask)}
def bounds(pred,pos,neg,universe):
    unknown=universe-pos-neg
    worst=pos|(unknown-pred);best=pos|(unknown&pred)
    lo=metric(pred,worst,universe);hi=metric(pred,best,universe)
    return {'universe_n':len(universe),'known_positive':len(pos),'known_negative':len(neg),'unknown_total':len(unknown),
        'unknown_predicted':len(unknown&pred),'unknown_not_predicted':len(unknown-pred),
        'f1_min':lo['f1'],'f1_max':hi['f1'],'precision_min':lo['precision'],'precision_max':hi['precision'],
        'recall_min':lo['recall'],'recall_max':hi['recall'],
        'all_unknown_negative_scenario':metric(pred,pos,universe),
        'all_unknown_positive_scenario':metric(pred,pos|unknown,universe)}
def self_test():
    u=set(range(7));p={0,2,3};pos={0,1};neg={2};unk=sorted(u-pos-neg)
    vals=[metric(p,pos|{v for v,b in zip(unk,bits) if b},u) for bits in itertools.product((0,1),repeat=len(unk))]
    b=bounds(p,pos,neg,u)
    for name in ('f1','precision','recall'):
        assert b[name+'_min']==min(v[name] for v in vals)
        assert b[name+'_max']==max(v[name] for v in vals)
    assert metric({0,2},{0,1},{0,1,2})['f1']==.5
    assert metric({0,2},{0,1},{0,1,2})['tp']==1
    assert metric({0,2},{0,1},{0,1,2})['fp']==1
    assert metric({0,2},{0,1},{0,1,2})['fn']==1
    assert metric({0},{0,1},{0,1})['recall']==.5
    return {'exhaustive_uncertainty_assignments_checked':len(vals),'metric_oracle_passed':True}
def main():
    tests=self_test();out=B/'analysis';out.mkdir(exist_ok=True)
    originals=json.loads((B/'ORIGINALS_MANIFEST.json').read_text(encoding='utf8'))
    pkg=json.loads(PKG.read_text(encoding='utf8'));frozen=json.loads(FROZEN.read_text(encoding='utf8'))
    assert sha(PRED)==frozen['predictions_sha256'] and pkg['version']==frozen['package_version']
    rawpred=json.loads(PRED.read_text(encoding='utf8'));ids=[s['id'] for s in pkg['sources']]
    assert ids==frozen['selected_source_ids'] and len(ids)==8
    legal={c+':'+i for c,items in pkg['taxonomy'].items() for i in items}
    pool={(s['id'],p['id']) for s in pkg['sources'] for p in s['candidates']}
    pred={method:{(sid,':'.join(pair)) for sid in ids for pair in rows[sid]} for method,rows in rawpred.items()}
    assert len(pool)==75 and set().union(*pred.values())==pool
    universe_all={(sid,pair) for sid in ids for pair in legal};assert len(universe_all)==1464
    reader_path=B/'code/excel_pilot.py'
    if not reader_path.exists():reader_path=ROOT/'tools/formal_mapping_scoring/excel_pilot.py'
    spec=importlib.util.spec_from_file_location('reader',reader_path);reader=importlib.util.module_from_spec(spec);spec.loader.exec_module(reader)
    records={};table=[];integrity=[];notes=[];firstpasses=[]
    for item in originals:
        original=B/item['relative_path'];assert sha(original)==item['sha256']
        copy=B/'working_copies'/item['role']/original.name;copy.parent.mkdir(parents=True,exist_ok=True)
        if copy.exists():assert sha(copy)==item['sha256']
        else:shutil.copyfile(original,copy)
        r=reader.read_workbook(copy,PKG);book=load_workbook(copy,data_only=False)
        labels={};missing={};formulas=[];non_answer_formula_headers=[]
        for row in range(9,84):
            s=book.worksheets[2];sid,pair=s.cell(row,1).value,s.cell(row,2).value
            key=(sid,pair);state=r['web_answer_fields'][sid]['candidates'][pair]
            assert key in pool and key not in labels;labels[key]=state
            table.append({'phase':item['role'].split('_')[0],'reviewer':item['role'][-1],'source_id':sid,'pair_id':pair,'kind':'candidate','state':state,'note':s.cell(row,6).value,'sheet':s.title,'label_cell':f'E{row}','reason_cell':f'F{row}'})
        for row in r['pair_rows'][reader.SHEETS[3]]:
            key=(row['source_id'],row['pair_id']);assert key not in pool and key not in missing;missing[key]=row['verdict']
            table.append({'phase':item['role'].split('_')[0],'reviewer':item['role'][-1],'source_id':key[0],'pair_id':key[1],'kind':'missing','state':row['verdict'],'note':row['note'],'sheet':reader.SHEETS[3],'label_cell':f"E{row['sheet_row']}",'reason_cell':f"F{row['sheet_row']}"})
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.data_type=='f' and ((sheet==book.worksheets[0] and cell.row<=139) or (sheet in [book.worksheets[1],book.worksheets[2],book.worksheets[3]] and cell.column in (1,2,5,6) and cell.row>=9)):formulas.append(sheet.title+'!'+cell.coordinate)
                    if cell.data_type=='f' and sheet==book.worksheets[0] and cell.row>139:non_answer_formula_headers.append({'cell':cell.coordinate,'raw':cell.value})
        assert not formulas
        for cell in ('B15','A143','A144'):
            value=book.worksheets[0][cell].value
            if value:notes.append({'role':item['role'],'sheet':book.worksheets[0].title,'cell':cell,'text':value})
        firstpasses.append({'role':item['role'],'rows':r['pair_rows'][reader.SHEETS[1]],'not_claimed_to_be_unexposed_historical_first_pass':True})
        records[item['role']]={'candidate':labels,'missing':missing,'all':labels|missing,'raw':r}
        integrity.append({'role':item['role'],'sha256':item['sha256'],'source_ids':ids,'package_version':r['package_version'],'reviewer_code':r['reviewer'],'template_slot':r['reviewer_slot'],'candidate_ids_exact':set(labels)==pool,'taxonomy_pairs':len(legal),'blank_candidates':sum(v=='unanswered' for v in labels.values()),'formulas_in_answer_fields':formulas,'non_answer_formula_headers':non_answer_formula_headers,'process_fields_kept_original':True,'full_reference_score_status':'blocked','full_reference_scores':None,'strict_reader_issues':r['issues']})
        book.close();assert sha(original)==item['sha256'] and sha(copy)==item['sha256']
    observed=pool|set().union(*(set(r['missing']) for r in records.values()));assert len(observed)==88
    summaries={};comparisons={};changes={};metrics={};sensitivity={};masks={}
    for role,r in records.items():
        summaries[role]={'candidate_counts':dict(Counter(r['candidate'].values())),'missing_counts':dict(Counter(r['missing'].values())),'first_pass_rows':len(r['raw']['pair_rows'][reader.SHEETS[1]]),'per_source':{sid:dict(Counter(v for (s,p),v in r['candidate'].items() if s==sid)) for sid in ids}}
        metrics[role]={};sensitivity[role]={}
        for name,domain in [('candidate_only',pool),('candidate_plus_observed_missing',observed)]:
            positive={k for k,v in r['all'].items() if v=='relevant'}&domain
            negative={k for k,v in r['all'].items() if v=='irrelevant'}&domain
            mask=positive|negative
            metrics[role][name]={'mask_n':len(mask),'domain_n':len(domain),'mask_fraction':len(mask)/len(domain),'positive_n':len(positive),'methods':{m:metric(p,positive,mask) for m,p in pred.items()}}
            masks[role+'_'+name]=sorted(mask)
        positive={k for k,v in r['all'].items() if v=='relevant'};negative={k for k,v in r['all'].items() if v=='irrelevant'}
        for name,domain in [('candidate_only',pool),('observed_union_88',observed),('full_taxonomy_1464_without_closed_world_assumption',universe_all)]:
            sensitivity[role][name]={m:bounds(p,positive&domain,negative&domain,domain) for m,p in pred.items()}
    common={}
    for phase in ('before','after'):
        c,d=records[phase+'_C'],records[phase+'_D'];cross=Counter((c['candidate'][k],d['candidate'][k]) for k in pool)
        same=[k for k in pool if c['candidate'][k]==d['candidate'][k]]
        shared=set(c['missing'])&set(d['missing'])
        comparisons[phase]={'same_candidate_labels':len(same),'same_fraction':len(same)/75,'both_uncertain':cross[('uncertain','uncertain')],
            'both_relevant':cross[('relevant','relevant')],'both_irrelevant':cross[('irrelevant','irrelevant')],
            'cross_table':[{ 'C':a,'D':b,'count':n} for (a,b),n in sorted(cross.items())],
            'candidate_disagreements':[{'source_id':k[0],'pair_id':k[1],'C':c['candidate'][k],'D':d['candidate'][k]} for k in sorted(pool-set(same))],
            'missing_union':len(set(c['missing'])|set(d['missing'])),'shared_missing':[{'source_id':k[0],'pair_id':k[1],'C':c['missing'][k],'D':d['missing'][k]} for k in sorted(shared)]}
        common[phase]={}
        for name,domain in [('candidate_only',pool),('candidate_plus_observed_missing',observed)]:
            mask={k for k in domain if c['all'].get(k) in ('relevant','irrelevant') and d['all'].get(k) in ('relevant','irrelevant')}
            common[phase][name]={'mask_n':len(mask),'domain_n':len(domain),'mask_fraction':len(mask)/len(domain),'source_mask_counts':{sid:sum(k[0]==sid for k in mask) for sid in ids},'reviewers':{}}
            masks['common_'+phase+'_'+name]=sorted(mask)
            for reviewer,r in [('C',c),('D',d)]:
                pos={k for k,v in r['all'].items() if v=='relevant'}
                common[phase][name]['reviewers'][reviewer]={m:{**metric(p,pos,mask),'per_source':[{'source_id':sid,**metric({k for k in p if k[0]==sid},{k for k in pos if k[0]==sid},{k for k in mask if k[0]==sid})} for sid in ids]} for m,p in pred.items()}
    for reviewer in ('C','D'):
        a,b=records['before_'+reviewer],records['after_'+reviewer]
        changes[reviewer]=[{'source_id':k[0],'pair_id':k[1],'before':a['candidate'][k],'after':b['candidate'][k]} for k in sorted(pool) if a['candidate'][k]!=b['candidate'][k]]
    assert summaries['before_C']['candidate_counts']=={'relevant':25,'irrelevant':31,'uncertain':19}
    assert summaries['before_D']['candidate_counts']=={'relevant':34,'irrelevant':26,'uncertain':15}
    assert summaries['after_D']['candidate_counts']=={'relevant':26,'irrelevant':31,'uncertain':18}
    assert comparisons['before']['same_candidate_labels']==59 and comparisons['after']['same_candidate_labels']==74
    assert len(changes['C'])==0 and len(changes['D'])==15
    assert common['after']['candidate_only']['mask_n']==56 and common['after']['candidate_plus_observed_missing']['mask_n']==57
    result={'purpose':'Post-hoc explicitly restricted descriptive statistics, not full formal gold/F1','full_gold_complete':False,'full_mapping_f1':None,
        'source_ids':ids,'candidate_n':75,'observed_union_n':88,'full_taxonomy_per_source':183,'taxonomy_frame_n':1464,
        'input_sha256':{'package':sha(PKG),'frozen_manifest':sha(FROZEN),'predictions':sha(PRED)},'originals':originals,
        'summaries':summaries,'comparisons':comparisons,'candidate_changes':changes,'per_reviewer_conditional_metrics':metrics,'common_definite_mask_metrics':common,
        'sensitivity':sensitivity,'tests':tests,
        'limitations':['All masks are post-hoc and selected by explicit label availability, never by method score. All 8 sources retained and per-source mask sizes disclosed.',
        'Candidate-only recall is conditional on a prediction-pooled universe; not complete mapping recall.',
        'Absence from one reviewer missing-pair list is unknown, not an explicit negative, even with full-taxonomy review confirmed.',
        'F1 extrema are method-specific possible completions, not one shared gold and not confidence intervals. All-unknown-negative/positive scenarios are shared completions, not actual labels.',
        'Post-discussion agreement is not independent inter-rater reliability. Saved before-C followed earlier feedback; do not assert pristine blind first-pass independence.',
        'User attestation confirms two human graduate-student reviewers and taxonomy coverage; original contradictory provenance notes are retained alongside it.',
        'Existing frozen full-reference protocol/scorer is unchanged; this is a newly authorized restricted supplemental analysis, not passing its original gate.']}
    save(out/'RESULTS.json',result);save(out/'INTEGRITY.json',integrity);save(out/'MASKS.json',masks);save(out/'ORIGINAL_NOTES.json',notes);save(out/'FIRST_PASS_RECORDS.json',firstpasses)
    with (out/'ANNOTATIONS_LONG.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(table[0]));w.writeheader();w.writerows(table)
    save(out/'ANNOTATIONS_LONG.json',table)
    save(out/'CANDIDATE_CHANGES.json',changes)
    print(json.dumps({'summaries':summaries,'comparisons':comparisons,'changes_n':{k:len(v) for k,v in changes.items()},'after_common':common['after'],'after_individual':{r:metrics[r] for r in ('after_C','after_D')},'after_sensitivity_observed':{r:sensitivity[r]['observed_union_88'] for r in ('after_C','after_D')}},ensure_ascii=True))
if __name__=='__main__':main()
