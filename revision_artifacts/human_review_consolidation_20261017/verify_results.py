"""Separate row-based verification; no calls to analysis metric functions."""
from pathlib import Path
import json,hashlib,itertools
from collections import Counter
from openpyxl import load_workbook
B=Path(__file__).parent
result=json.loads((B/'analysis/RESULTS.json').read_text(encoding='utf8'))
manifest=json.loads((B/'ORIGINALS_MANIFEST.json').read_text(encoding='utf8'))
data={};checks=[]
for f in manifest:
    p=B/f['relative_path'];assert hashlib.sha256(p.read_bytes()).hexdigest()==f['sha256']
    w=load_workbook(p,data_only=False);rows={}
    for i in range(9,84):
        s=w.worksheets[2];key=(s.cell(i,1).value,s.cell(i,2).value)
        assert key not in rows;rows[key]=s.cell(i,5).value
    assert len(rows)==75
    for i in range(9,169):
        s=w.worksheets[3];key=(s.cell(i,1).value,s.cell(i,2).value)
        if key[0] is not None:
            assert key not in rows;rows[key]=s.cell(i,5).value
    data[f['role']]=rows;w.close()
checks.append('Four original SHA256 and stable-key uniqueness verified from XLSX')
pool={(x['source_id'],x['pair_id']) for x in json.loads((B/'analysis/ANNOTATIONS_LONG.json').read_text(encoding='utf8')) if x['kind']=='candidate'}
for phase,count in [('before',59),('after',74)]:
    assert sum(data[phase+'_C'][k]==data[phase+'_D'][k] for k in pool)==count
assert sum(data['before_D'][k]!=data['after_D'][k] for k in pool)==15
checks.append('Agreement and actual changes independently recounted')
common={k for k in data['after_C'] if data['after_C'][k] in ('相关','不相关') and data['after_D'].get(k) in ('相关','不相关')}
assert len(common)==57 and all(data['after_C'][k]==data['after_D'][k] for k in common)
predpath=B/'inputs/e1_predictions.json'
if not predpath.exists():predpath=Path(r'D:\cooling_rag_multigranular\cooling_rag_multigranular\data\search_results\exp_v4\e1_predictions.json')
pred=json.loads(predpath.read_text(encoding='utf8'))
oracle={'deepseek':(21,1,5,30),'qwen':(18,3,8,28),'bm25':(6,20,20,11),'embed':(12,11,14,20)}
for method in oracle:
    counts=Counter()
    for sid,pair in common:
        hit=pair.split(':') in pred[method][sid]
        gold=data['after_C'][(sid,pair)]=='相关'
        counts['tp' if hit and gold else 'fp' if hit else 'fn' if gold else 'tn']+=1
    actual=tuple(counts[k] for k in ('tp','fp','fn','tn'));assert actual==oracle[method]
    r=result['common_definite_mask_metrics']['after']['candidate_plus_observed_missing']['reviewers']['C'][method]
    assert actual==tuple(r[k] for k in ('tp','fp','fn','tn'))
    assert sum(actual)==57 and sum(x['mask_n'] for x in r['per_source'])==57
checks.append('Independent rowwise TP/FP/FN/TN matched 4 hand-checkable count oracles on common 57 mask')
strict=json.loads((B/'analysis/STRICT_FULL_REFERENCE_CHECK.json').read_text(encoding='utf8'))
assert strict['status']=='blocked' and strict['scores'] is None and result['full_mapping_f1'] is None
checks.append('Original full-reference gate remains blocked with scores=null; subset statistics clearly separate')
report={'passed':True,'checks':checks,'source_ids':result['source_ids'],'counts_oracle':oracle,'uncertainty_formula_test':result['tests']}
(B/'analysis/INDEPENDENT_VERIFICATION.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=True))
