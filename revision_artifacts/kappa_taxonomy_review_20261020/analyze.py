"""Frozen analysis of the blind taxonomy-label review.

Usage: python analyze.py returned.xlsx --out review_results/reviewer1.json
Blocks if blind ids or displayed text were changed, or any primary label is missing.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import random
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
COORD = HERE / "delivery/02_协调者_勿发"
UNKNOWN = "无法判断"
SEED, B = 20261020, 10000


def code(v) -> str | None:
    if v is None or str(v).strip() == "":
        return None
    s = str(v).strip()
    return UNKNOWN if s == UNKNOWN else s.split("｜")[0].strip()


def kappa(a: list, b: list) -> float | None:
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return None if pe == 1 else (po - pe) / (1 - pe)


def boot(a: list, b: list, fn) -> list[float]:
    rng = random.Random(SEED); n = len(a); vals = []
    for _ in range(B):
        idx = [rng.randrange(n) for _ in range(n)]
        v = fn([a[i] for i in idx], [b[i] for i in idx])
        if v is not None:
            vals.append(v)
    vals.sort()
    return [vals[int(0.025 * len(vals))], vals[int(0.975 * len(vals)) - 1]]


def agree(a, b):
    return sum(x == y for x, y in zip(a, b)) / len(a)


def read(path: Path, key: list[dict]) -> tuple[list[dict], list[str], dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    ws = wb["标注50"]; issues = []; out = []
    for r in key:
        row = [c.value for c in ws[r["row"] + 1]][:10]
        if row[0] != r["blind_id"] or (row[3] or "").strip() != r["display_text"]:
            issues.append(f"{r['blind_id']}: id or text changed")
        lab = {"comp": code(row[4]), "int": code(row[5]), "alt_comp": code(row[6]), "alt_int": code(row[7]),
               "confidence": row[8], "note": row[9]}
        if lab["comp"] is None or lab["int"] is None:
            issues.append(f"{r['blind_id']}: primary label missing")
        out.append({**r, **{f"rev_{k}": v for k, v in lab.items()}})
    ident = {}
    s0 = wb["说明"]
    for row in s0.iter_rows(values_only=True):
        if row and row[0] and len(row) > 1 and row[1] not in (None, ""):
            ident[str(row[0])] = row[1]
    wb.close()
    return out, issues, ident


def summarize(rows: list[dict]) -> dict:
    ac = [r["author_component"] for r in rows]; rc = [r["rev_comp"] for r in rows]
    ai = [r["author_intent"] for r in rows]; ri = [r["rev_int"] for r in rows]
    ap = [f"{x}/{y}" for x, y in zip(ac, ai)]; rp = [f"{x}/{y}" for x, y in zip(rc, ri)]
    lenient = [r["author_component"] == r["rev_comp"] and r["author_intent"] == r["rev_int"]
               or (r["author_component"] == r["rev_alt_comp"] and r["author_intent"] == r["rev_alt_int"]) for r in rows]
    res = {"n": len(rows),
           "component": {"agreement": agree(ac, rc), "kappa": kappa(ac, rc), "kappa_ci95": boot(ac, rc, kappa)},
           "aspect": {"agreement": agree(ai, ri), "kappa": kappa(ai, ri), "kappa_ci95": boot(ai, ri, kappa)},
           "pair": {"agreement": agree(ap, rp), "kappa": kappa(ap, rp), "agreement_ci95": boot(ap, rp, agree),
                    "lenient_agreement": sum(lenient) / len(rows)},
           "reviewer_unknown": {"component": rc.count(UNKNOWN), "aspect": ri.count(UNKNOWN)}}
    known = [i for i in range(len(rows)) if rc[i] != UNKNOWN and ri[i] != UNKNOWN]
    if known:
        res["excluding_unknown"] = {"n": len(known),
                                    "component_kappa": kappa([ac[i] for i in known], [rc[i] for i in known]),
                                    "aspect_kappa": kappa([ai[i] for i in known], [ri[i] for i in known]),
                                    "pair_agreement": agree([ap[i] for i in known], [rp[i] for i in known])}
    res["by_tier"] = {}
    for t in sorted({r["kb_tier"] for r in rows}):
        idx = [i for i, r in enumerate(rows) if r["kb_tier"] == t]
        res["by_tier"][t] = {"n": len(idx), "component_agreement": agree([ac[i] for i in idx], [rc[i] for i in idx]),
                             "aspect_agreement": agree([ai[i] for i in idx], [ri[i] for i in idx]),
                             "pair_agreement": agree([ap[i] for i in idx], [rp[i] for i in idx])}
    res["author_pairs_outside_183"] = sum(not r["author_pair_valid"] for r in rows)
    res["disagreements"] = [{k: r[k] for k in ("blind_id", "unit_id", "kb_tier", "author_component", "author_intent",
                                                "rev_comp", "rev_int", "rev_alt_comp", "rev_alt_int", "rev_note")}
                            for r, ok in zip(rows, lenient) if not ok]
    return res


def main() -> int:
    ap = argparse.ArgumentParser(); ap.add_argument("workbook", type=Path); ap.add_argument("--out", type=Path)
    a = ap.parse_args()
    proto = json.loads((COORD / "PROTOCOL_FROZEN.json").read_text(encoding="utf-8"))
    keyp = COORD / "SAMPLE_KEY.json"
    if hashlib.sha256(keyp.read_bytes()).hexdigest() != proto["sha256"]["key"]:
        raise SystemExit("frozen key changed")
    if hashlib.sha256(Path(__file__).read_bytes()).hexdigest() != proto["sha256"]["analyze.py"]:
        raise SystemExit("frozen analysis changed")
    key = json.loads(keyp.read_text(encoding="utf-8"))
    rows, issues, ident = read(a.workbook, key)
    if issues:
        print(json.dumps({"status": "blocked", "issues": issues}, ensure_ascii=False, indent=2)); return 2
    result = {"status": "complete", "input_sha256": hashlib.sha256(a.workbook.read_bytes()).hexdigest(),
              "identity": ident, "summary": summarize(rows),
              "raw": [{k: r[k] for k in r if k != "display_text"} for r in rows]}
    text = json.dumps(result, ensure_ascii=False, indent=2, default=str)
    if a.out:
        if a.out.exists():
            raise SystemExit("refusing overwrite")
        a.out.parent.mkdir(parents=True, exist_ok=True); a.out.write_text(text + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in result["summary"].items() if k != "disagreements"}, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
