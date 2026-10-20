"""Synthetic tests for build.py / analyze.py (no real ratings)."""
import importlib.util
import json
import random
from pathlib import Path

import openpyxl
import pytest

HERE = Path(__file__).resolve().parent


def load(name):
    spec = importlib.util.spec_from_file_location(name, HERE / f"{name}.py")
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m


bld, ana = load("build"), load("analyze")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    d = tmp_path_factory.mktemp("k")
    units = [json.loads(l) for l in bld.KB.open(encoding="utf-8")]
    rows = bld.sample(units)
    wb = d / "blank.xlsx"; bld.build_workbook(rows, wb)
    return rows, wb, d


def fill(wb, rows, out, fn):
    w = openpyxl.load_workbook(wb); ws = w["标注50"]
    for r in rows:
        c, i = fn(r)
        ws.cell(r["row"] + 1, 5, c); ws.cell(r["row"] + 1, 6, i)
    w.save(out); return out


def test_sample_deterministic_and_text_clean(built):
    rows, _, _ = built
    units = [json.loads(l) for l in bld.KB.open(encoding="utf-8")]
    assert [r["unit_id"] for r in bld.sample(units)] == [r["unit_id"] for r in rows]
    assert len({r["unit_id"] for r in rows}) == 50
    for r in rows:
        assert "整理说明" not in r["display_text"] and "可支撑" not in r["display_text"]
        assert r["display_text"]


def test_perfect_agreement_kappa_one(built):
    rows, wb, d = built
    p = fill(wb, rows, d / "perfect.xlsx", lambda r: (bld.option(r["author_component"], bld.COMPONENTS[r["author_component"]]),
                                                     bld.option(r["author_intent"], bld.ASPECTS[r["author_intent"]])))
    got, issues, _ = ana.read(p, rows)
    assert not issues
    s = ana.summarize(got)
    assert s["component"]["kappa"] == pytest.approx(1) and s["aspect"]["kappa"] == pytest.approx(1)
    assert s["pair"]["agreement"] == 1 and s["disagreements"] == []


def test_random_labels_low_kappa_and_unknown(built):
    rows, wb, d = built
    rng = random.Random(1); cs, its = list(bld.COMPONENTS), list(bld.ASPECTS)
    p = fill(wb, rows, d / "rand.xlsx", lambda r: (rng.choice(cs + [bld.UNKNOWN]), rng.choice(its)))
    got, issues, _ = ana.read(p, rows)
    assert not issues
    s = ana.summarize(got)
    assert s["component"]["kappa"] < 0.3 and s["aspect"]["kappa"] < 0.3


def test_blocks_missing_and_changed(built):
    rows, wb, d = built
    p = fill(wb, rows, d / "miss.xlsx", lambda r: (None, None))
    _, issues, _ = ana.read(p, rows)
    assert any("missing" in i for i in issues)
    w = openpyxl.load_workbook(p); w["标注50"].cell(2, 4, "篡改"); w.save(d / "chg.xlsx")
    _, issues, _ = ana.read(d / "chg.xlsx", rows)
    assert any("changed" in i for i in issues)


def test_kappa_known_value():
    a = ["x"] * 5 + ["y"] * 5; b = ["x"] * 4 + ["y"] * 6
    assert ana.kappa(a, b) == pytest.approx(0.8)
