"""Build the blind taxonomy-label review (R1-4.2): 50 random KB units, external reviewer.

The reviewer assigns one primary (component, aspect) to each unit WITHOUT seeing the
authors' label; analyze.py then computes Cohen's kappa against the KB labels.

Display text: the unit text up to the first curator line (整理说明 / 可支撑 / 不能支撑),
which state the assigned label in prose; source document and location are kept.
Writes delivery/01_发给评审/ (workbook only) and delivery/02_协调者_勿发/ (key, protocol, hashes).
"""
from __future__ import annotations

import hashlib
import json
import random
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
from cooling_rag_mg.evidence_taxonomy import COMPONENTS, INTENTS, TAXONOMY, valid_pair  # noqa: E402

KB = ROOT / "data/processed/evidence_units_v3_audited.jsonl"
OUT = HERE / "delivery"
SEND, COORD = OUT / "01_发给评审", OUT / "02_协调者_勿发"
SEED, N = 20261020, 50
CURATOR = re.compile(r"^\s*(整理说明|可支撑|不能支撑)")
UNKNOWN = "无法判断"
sha = lambda b: hashlib.sha256(b).hexdigest()


def display_text(text: str) -> str:
    keep = []
    for line in text.split("\n"):
        if CURATOR.match(line):
            break
        keep.append(line)
    return "\n".join(keep).strip()


def option(code: str, name: str) -> str:
    return f"{code}｜{name}"


def aspect_names() -> dict[str, str]:
    """INTENTS plus the component-specific aspects used in TAXONOMY but defined only by an
    inline comment in evidence_taxonomy.py (first comment wins)."""
    names = dict(INTENTS)
    src = (ROOT / "cooling_rag_mg/evidence_taxonomy.py").read_text(encoding="utf-8")
    for code_, comment in re.findall(r'^\s*"([a-z_]+)",\s*#\s*(.+)$', src, re.M):
        names.setdefault(code_, comment.strip())
    used = {i for v in TAXONOMY.values() for i in v}
    missing = used - set(names)
    if missing:
        raise SystemExit(f"aspects without a description: {sorted(missing)}")
    return names


ASPECTS = aspect_names()


def sample(units: list[dict]) -> list[dict]:
    ids = sorted(u["id"] for u in units)
    chosen = random.Random(SEED).sample(ids, N)
    by = {u["id"]: u for u in units}
    rows = []
    for k, uid in enumerate(chosen, 1):
        u = by[uid]
        rows.append({"row": k, "blind_id": f"T{k:02d}", "unit_id": uid, "kb_tier": u["kb_tier"],
                     "author_component": u["primary_component"], "author_intent": u["primary_intent"],
                     "author_pair_valid": valid_pair(u["primary_component"], u["primary_intent"]),
                     "source_document": u["source_document"], "location": u["location"],
                     "display_text": display_text(u["text"])})
    return rows


INSTRUCTIONS = [
    "证据单元分类标注（独立盲评）",
    "",
    "任务：为每条证据单元选择一个最主要的“部件”和一个最主要的“工程方面”。",
    "1. 只依据该条原文判断。若原文涉及多个部件或方面，选最主要的一个；可在“备选”列再给一组。",
    "2. 部件、方面均从下拉框选择（选项见“选项”页，含义见“定义”页）。确实无法判断时选“无法判断”。",
    "3. “有效组合”页列出183个常用（部件，方面）组合，供参考；若你认为正确的组合不在表中，仍按你的判断选。",
    "4. 请独立完成，不与他人讨论，不使用AI给出分类建议；可查阅原始文件。",
    "5. 只填黄色列；不要修改编号和原文。预计约1–1.5小时。",
]
IDENTITY = ["评审代号", "身份（外部非作者 / 作者 / 其他）", "相关专业", "是否独立完成且未使用AI分类建议（是/否）", "备注"]


def build_workbook(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws0 = wb.active; ws0.title = "说明"
    for i, line in enumerate(INSTRUCTIONS, 1):
        ws0.cell(i, 1, line).font = Font(bold=(i == 1), size=12 if i == 1 else 11)
    base = len(INSTRUCTIONS) + 2
    ws0.cell(base, 1, "评审人填写").font = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for j, label in enumerate(IDENTITY, 1):
        ws0.cell(base + j, 1, label); ws0.cell(base + j, 2).fill = yellow
    ws0.column_dimensions["A"].width = 46; ws0.column_dimensions["B"].width = 60

    opt = wb.create_sheet("选项")
    comps = [option(c, n) for c, n in COMPONENTS.items()] + [UNKNOWN]
    ints = [option(c, n) for c, n in ASPECTS.items()] + [UNKNOWN]
    opt.cell(1, 1, "部件"); opt.cell(1, 2, "工程方面")
    for i, v in enumerate(comps, 2):
        opt.cell(i, 1, v)
    for i, v in enumerate(ints, 2):
        opt.cell(i, 2, v)
    opt.column_dimensions["A"].width = 50; opt.column_dimensions["B"].width = 70

    ws = wb.create_sheet("标注50", 1)
    head = ["编号", "来源文件", "位置", "证据原文", "主部件", "主工程方面", "备选部件（可选）", "备选工程方面（可选）", "把握程度1–3", "备注"]
    for j, h in enumerate(head, 1):
        ws.cell(1, j, h).font = Font(bold=True)
    for i, r in enumerate(rows, 2):
        for j, v in enumerate([r["blind_id"], r["source_document"], r["location"], r["display_text"]], 1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
        for j in range(5, 11):
            ws.cell(i, j).fill = yellow
    last = len(rows) + 1
    dv_c = DataValidation(type="list", formula1=f"='选项'!$A$2:$A${len(comps) + 1}", allow_blank=True)
    dv_i = DataValidation(type="list", formula1=f"='选项'!$B$2:$B${len(ints) + 1}", allow_blank=True)
    dv_n = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    for dv in (dv_c, dv_i, dv_n):
        ws.add_data_validation(dv)
    dv_c.add(f"E2:E{last}"); dv_c.add(f"G2:G{last}")
    dv_i.add(f"F2:F{last}"); dv_i.add(f"H2:H{last}")
    dv_n.add(f"I2:I{last}")
    for col, w in zip("ABCDEFGHIJ", (7, 28, 10, 80, 30, 34, 26, 30, 10, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "E2"

    pairs = wb.create_sheet("有效组合")
    pairs.append(["部件", "工程方面"])
    for c, intents in TAXONOMY.items():
        for it in intents:
            pairs.append([option(c, COMPONENTS[c]), option(it, ASPECTS[it])])
    pairs.column_dimensions["A"].width = 50; pairs.column_dimensions["B"].width = 70

    defs = wb.create_sheet("定义")
    defs.append(["类别", "代码", "含义"])
    for c, n in COMPONENTS.items():
        defs.append(["部件", c, n])
    for c, n in ASPECTS.items():
        defs.append(["工程方面", c, n])
    defs.column_dimensions["B"].width = 30; defs.column_dimensions["C"].width = 80
    wb.properties.creator = None
    wb.save(path)


def main() -> None:
    if OUT.exists():
        raise SystemExit("delivery exists; refusing to rebuild a frozen sample")
    units = [json.loads(l) for l in KB.open(encoding="utf-8")]
    rows = sample(units)
    SEND.mkdir(parents=True); COORD.mkdir(parents=True)
    wbp = SEND / "证据单元分类标注50.xlsx"
    build_workbook(rows, wbp)
    key = COORD / "SAMPLE_KEY.json"
    key.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    protocol = {
        "version": "taxonomy-kappa-v1", "frozen_utc": datetime.now(timezone.utc).isoformat(),
        "kb_sha256": sha(KB.read_bytes()), "kb_units": len(units), "seed": SEED, "n": N,
        "sampling": "simple random sample of unit ids (sorted) with random.Random(seed).sample",
        "display": "unit text up to the first curator line (整理说明/可支撑/不能支撑); source document and location shown; author label hidden",
        "primary": "Cohen's kappa between reviewer primary label and KB label, separately for component and aspect; "
                   "无法判断 counted as disagreement (its own category)",
        "secondary": ["exact (component, aspect) pair agreement and kappa over pairs",
                      "lenient pair agreement: KB pair equals reviewer primary or alternative pair",
                      "kappa excluding 无法判断", "agreement by KB tier"],
        "interval": "unit bootstrap, 10000 resamples, seed 20261020, percentile 95%",
        "author_pairs_outside_183": sum(not r["author_pair_valid"] for r in rows),
        "sha256": {"workbook": sha(wbp.read_bytes()), "key": sha(key.read_bytes()),
                   "build.py": sha((HERE / "build.py").read_bytes()), "analyze.py": sha((HERE / "analyze.py").read_bytes())},
    }
    (COORD / "PROTOCOL_FROZEN.json").write_text(json.dumps(protocol, ensure_ascii=False, indent=2), encoding="utf-8")
    tiers = {}
    for r in rows:
        tiers[r["kb_tier"]] = tiers.get(r["kb_tier"], 0) + 1
    print(json.dumps({"n": len(rows), "tiers": tiers, "outside_183": protocol["author_pairs_outside_183"],
                      "max_text": max(len(r["display_text"]) for r in rows), "workbook": str(wbp)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
