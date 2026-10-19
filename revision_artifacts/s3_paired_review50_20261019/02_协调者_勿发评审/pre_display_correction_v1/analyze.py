"""Frozen local analysis. Rejects blank, altered, incomplete or non-first-pass forms."""
import argparse, hashlib, json, math, random, sys
from collections import Counter, defaultdict
from datetime import datetime
from fractions import Fraction
from pathlib import Path
import openpyxl

HERE=Path(__file__).resolve().parent
COORD=HERE/'02_协调者_勿发评审'
CORRECTNESS=['正确','可接受','错误','无法判断']
APPLICABILITY=['适用','不适用','工况不明']
ROLES=['外部非作者领域评审','共同作者领域评审','其他（备注说明）']
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
read=lambda p:json.loads(Path(p).read_text(encoding='utf8'))
def norm(v):return None if v in (None,'') else v

def allowed(sheet,row,col):
    return (sheet=='身份与说明' and col==2 and 6<=row<=14) or (sheet=='评分50' and 6<=row<=55 and 2<=col<=6)

def contract(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
    result={}
    for ws in wb:
        cells={}
        for row in ws.iter_rows():
            for c in row:
                if norm(c.value) is not None and not allowed(ws.title,c.row,c.column):
                    if c.data_type=='f':raise ValueError('Unexpected formula in template')
                    cells[c.coordinate]=c.value
        result[ws.title]={'max_row':ws.max_row,'max_column':ws.max_column,'cells':cells}
    wb.close();return result

def validate(path,expected,sample):
    issues=[];ratings=[];metadata={}
    wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
    if set(wb.sheetnames)!=set(expected):issues.append('工作表集合改变')
    for name,ref in expected.items():
        if name not in wb.sheetnames:continue
        ws=wb[name]
        if ws.max_row!=ref['max_row'] or ws.max_column!=ref['max_column']:issues.append(f'{name}:行列范围改变')
        observed={}
        for row in ws.iter_rows():
            for c in row:
                if c.data_type=='f':issues.append(f'{name}!{c.coordinate}:禁止公式答案或原文公式')
                if norm(c.value) is not None and not allowed(name,c.row,c.column):observed[c.coordinate]=c.value
        if observed!=ref['cells']:issues.append(f'{name}:原文/盲ID/定义/结构被修改')
    if '身份与说明' in wb.sheetnames:
        ws=wb['身份与说明']
        metadata={k:norm(ws[f'B{r}'].value) for k,r in [('reviewer_id',6),('role',7),('expertise',8),('organization',9),('prior_exposure',10),('independent_first_pass',11),('record_stage',12),('start_time',13),('end_time',14)]}
        for k,v in metadata.items():
            if k!='organization' and v is None:issues.append('身份/过程字段未填:'+k)
        if metadata['role'] not in ROLES:issues.append('身份选项无效')
        if metadata['independent_first_pass']!='是' or metadata['record_stage']!='独立首判':issues.append('仅分析明确声明独立完成的首判；合议或不独立表不得算独立一致性')
        parsed=[]
        for k in ['start_time','end_time']:
            v=metadata[k]
            try:
                dt=v if isinstance(v,datetime) else datetime.fromisoformat(str(v).replace('Z','+00:00'))
                if dt.tzinfo is None:raise ValueError('timezone missing')
                parsed.append(dt)
            except (TypeError,ValueError):issues.append(k+':请写含时区ISO时间，例如2026-10-20T09:00:00+08:00（例子不是填写值）')
        if len(parsed)==2 and parsed[1]<parsed[0]:issues.append('完成时间早于开始时间')
    byid={r['blind_id']:r for r in sample}
    if '评分50' in wb.sheetnames:
        ws=wb['评分50'];seen=[]
        for i in range(6,56):
            code,c,a,basis,confidence,note=[norm(ws.cell(i,j).value) for j in range(1,7)]
            seen.append(code)
            if code not in byid:issues.append(f'未知盲ID:{code}');continue
            if c not in CORRECTNESS:issues.append(f'{code}:正确性未填或非法')
            if a not in APPLICABILITY:issues.append(f'{code}:适用性未填或非法')
            if str(confidence) not in ['1','2','3']:issues.append(f'{code}:信心须为1/2/3')
            if c in ['正确','可接受'] and (basis is None or not str(basis).strip()):issues.append(f'{code}:正确/可接受必须填写具体依据')
            ratings.append({'blind_id':code,'source_id':byid[code]['source_id'],'arm':byid[code]['arm'],'correctness':c,'applicability':a,'basis':basis,'confidence':confidence,'note':note})
        if len(set(seen))!=50 or set(seen)!=set(byid):issues.append('盲ID重复/缺失')
    wb.close()
    return {'valid':not issues,'issues':issues,'metadata':metadata,'ratings':ratings if not issues else None,'input_sha256':sha(path)}

def status(r,secondary=False):
    labels=['正确','可接受'] if secondary else ['正确']
    c,a=r['correctness'],r['applicability']
    success=c in labels and a=='适用'
    unresolved=(c in labels or c=='无法判断') and a!='不适用' and (c=='无法判断' or a=='工况不明')
    return int(success),int(unresolved)

def sign_test(diffs):
    pos=sum(x>0 for x in diffs);neg=sum(x<0 for x in diffs);ties=len(diffs)-pos-neg;n=pos+neg
    p=None if not n else min(1.,2*sum(math.comb(n,k) for k in range(min(pos,neg)+1))/2**n)
    return {'positive_sources':pos,'negative_sources':neg,'ties':ties,'n_nonzero':n,'two_sided_p':p,'scope':'descriptive exact sign test; assumes independent source signs and equal positive/negative probability under null; no equivalence/power claim'}

def percentile(values,q):
    vs=sorted(values);z=(len(vs)-1)*q;i=int(z);j=min(i+1,len(vs)-1);return vs[i]+(vs[j]-vs[i])*(z-i)

def outcome(rows,secondary=False):
    grouped=defaultdict(list)
    for r in rows:grouped[(r['source_id'],r['arm'])].append(r)
    sources=sorted({s for s,a in grouped});src=[];diffs=[];lo=[];hi=[]
    arm_source=defaultdict(list)
    for s in sources:
        rates={}
        for a in ['rag','norag']:
            rs=grouped[s,a]
            if not rs:raise ValueError('Unpaired source')
            pairs=[status(r,secondary) for r in rs];lower=Fraction(sum(t[0] for t in pairs),len(rs));upper=lower+Fraction(sum(t[1] for t in pairs),len(rs))
            rates[a]={'n':len(rs),'confirmed_success':float(lower),'unresolved_possible_success':float(upper-lower),'possible_success_upper':float(upper)};arm_source[a].append(lower)
            if a=='rag':rl,ru=lower,upper
            else:nl,nu=lower,upper
        diff=rl-nl;diffs.append(diff);lo.append(rl-nu);hi.append(ru-nl);src.append({'source_id':s,**rates,'confirmed_success_difference':float(diff),'difference_exact':str(diff),'unfavorable_bound':float(lo[-1]),'favorable_bound':float(hi[-1])})
    rng=random.Random(202610191556);floats=list(map(float,diffs));boot=[sum(rng.choices(floats,k=len(floats)))/len(floats) for _ in range(10000)]
    item={}
    for a in ['rag','norag']:
        rs=[r for r in rows if r['arm']==a];item[a]={'n':len(rs),'confirmed_success_count':sum(status(r,secondary)[0] for r in rs),'confirmed_success_fraction':sum(status(r,secondary)[0] for r in rs)/len(rs),'unresolved_count':sum(status(r,secondary)[1] for r in rs)}
    return {'source_count':len(sources),'source_equal_confirmed_success':{a:float(sum(v)/len(v)) for a,v in arm_source.items()},'source_equal_difference':float(sum(diffs)/len(diffs)),'source_equal_difference_exact':str(sum(diffs)/len(diffs)),'unknown_assignment_bounds':[float(sum(lo)/len(lo)),float(sum(hi)/len(hi))],'source_bootstrap_percentile_interval':[percentile(boot,.025),percentile(boot,.975)],'interval_scope':'exploratory 10000 source resamples; seed202610191556; 20 eligible common sources, not all60; not guaranteed population CI and not model run variability','sign_test':sign_test(diffs) if not secondary else None,'per_source':src,'item_descriptive':item}

def summarize(rows):
    descriptions={}
    for a in ['rag','norag']:
        rs=[r for r in rows if r['arm']==a];sources=sorted({r['source_id'] for r in rs})
        functions={'error':lambda r:r['correctness']=='错误','unknown':lambda r:r['correctness']=='无法判断' or r['applicability']=='工况不明','not_applicable':lambda r:r['applicability']=='不适用'}
        descriptions[a]={'n':len(rs),'correctness':dict(Counter(r['correctness'] for r in rs)),'applicability':dict(Counter(r['applicability'] for r in rs)),'practice_basis_items':sum('实践判断' in str(r['basis']) for r in rs)}
        for label,fn in functions.items():
            descriptions[a][label+'_item_rate']=sum(fn(r) for r in rs)/len(rs)
            descriptions[a][label+'_source_equal_rate']=sum(sum(fn(r) for r in rs if r['source_id']==s)/sum(r['source_id']==s for r in rs) for s in sources)/len(sources)
    return {'primary':outcome(rows),'secondary':outcome(rows,True),'labels_and_risks':descriptions,'interpretation':'Confirmed-success rate counts observed positive joint judgments; unknown is separately retained, not imputed factual error. Bounds express unresolved outcomes. Reviewers analyzed separately; no pooled gold.'}

def agreement(first,second):
    b={r['blind_id']:r for r in second};out={}
    for field,labels in [('correctness',CORRECTNESS),('applicability',APPLICABILITY)]:
        matrix={x:{y:0 for y in labels} for x in labels}
        for r in first:matrix[r[field]][b[r['blind_id']][field]]+=1
        n=len(first);po=Fraction(sum(matrix[x][x] for x in labels),n)
        pe=sum(Fraction(sum(matrix[x].values())*sum(matrix[y][x] for y in labels),n*n) for x in labels)
        out[field]={'n':n,'exact_agreement':float(po),'kappa':None if pe==1 else float((po-pe)/(1-pe)),'matrix_rows_first_columns_second':matrix,'unknown_retained_as_nominal_category':True,'interpretation':'descriptive only; independent-first-pass status self-declared; shared evidence or coauthor involvement limits external independence'}
    return out

def main():
    parser=argparse.ArgumentParser();parser.add_argument('workbooks',nargs='+',type=Path);parser.add_argument('--out',type=Path);args=parser.parse_args()
    if not 1<=len(args.workbooks)<=2:raise SystemExit('Provide one or two first-pass workbooks')
    freeze=read(COORD/'ANALYSIS_FREEZE.json')
    for name,h in freeze['files'].items():
        if sha(HERE/name)!=h:raise SystemExit('Frozen analysis/input changed: '+name)
    sample=read(COORD/'SAMPLE_FROZEN.json');expected=read(COORD/'WORKBOOK_CONTRACT.json')
    validations=[validate(p,expected,sample) for p in args.workbooks]
    if not all(v['valid'] for v in validations):
        print(json.dumps({'status':'blocked','scores':None,'validation':[{k:v[k] for k in ['valid','issues','input_sha256']} for v in validations]},ensure_ascii=False,indent=2));return 2
    if len(validations)==2 and str(validations[0]['metadata']['reviewer_id']).strip().casefold()==str(validations[1]['metadata']['reviewer_id']).strip().casefold():raise SystemExit('Distinct reviewer identities required; no agreement computed')
    result={'status':'complete_first_pass_forms','external_non_author_present':any(v['metadata']['role']=='外部非作者领域评审' for v in validations),'reviewers':[{'metadata':v['metadata'],'input_sha256':v['input_sha256'],'raw_ratings':v['ratings'],'summary':summarize(v['ratings'])} for v in validations],'agreement':agreement(validations[0]['ratings'],validations[1]['ratings']) if len(validations)==2 else None,'caution':'No external status inferred from author identity; declarations not independently authenticated. No significant difference is not equivalence.'}
    if not result['external_non_author_present']:result['remaining_limit']='External non-author domain assessment not yet present; this is not independent external validation.'
    content=json.dumps(result,ensure_ascii=False,indent=2,default=str)
    if args.out:
        out=args.out.resolve();base=(HERE/'review_results').resolve()
        if not out.is_relative_to(base) or out.exists():raise SystemExit('Output must be a new file under this task/review_results; refusing overwrite')
        out.parent.mkdir(parents=True,exist_ok=True);out.write_text(content+'\n',encoding='utf8')
    print(content);return 0
if __name__=='__main__':sys.exit(main())
