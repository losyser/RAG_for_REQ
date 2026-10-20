"""Preserve attached workbook and run unchanged frozen validator; never infer answers."""
import json,hashlib,subprocess,sys
from pathlib import Path
from collections import Counter
from datetime import datetime,timezone
import openpyxl

H=Path(__file__).resolve().parent;ROOT=H.parents[1];FROZEN=ROOT/'revision_artifacts/s3_paired_review50_20261019';COORD=FROZEN/'02_协调者_勿发评审'
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
read=lambda p:json.loads(Path(p).read_text(encoding='utf8'))
def dump(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf8')

def main():
    raw=H/'raw/数值候选独立评分.xlsx';expected_hash='87f3153b7ddb329efd528a45621135a812298c0ca812d42eb59caae6a03a173f'
    assert raw.stat().st_size==35725 and sha(raw)==expected_hash
    protocol=read(COORD/'ANALYSIS_FREEZE.json');checks={name:sha(FROZEN/name)==h for name,h in protocol['files'].items()};assert all(checks.values())
    protected_paths=read(ROOT/'revision_artifacts/s3_reference_supplement_20261020/02_工作记录_不发评审/PROTECTED_BEFORE.json')
    current={p:sha(p) for p in protected_paths};dump(H/'PROTECTED_AT_RETURN_START.json',current)
    before=sha(raw)
    result=subprocess.run([sys.executable,'-E','-B','-X','utf8',str(FROZEN/'analyze.py'),str(raw)],capture_output=True,text=True,encoding='utf8')
    (H/'FROZEN_ANALYSIS_STDOUT.json').write_text(result.stdout,encoding='utf8');(H/'FROZEN_ANALYSIS_STDERR.txt').write_text(result.stderr,encoding='utf8')
    response=json.loads(result.stdout);assert result.returncode==2 and response['status']=='blocked' and response['scores'] is None
    w=openpyxl.load_workbook(raw,read_only=True,data_only=False);contract=read(COORD/'WORKBOOK_CONTRACT.json')
    sheet_checks={};modified=[]
    for name,ref in contract.items():
        ws=w[name];actual={}
        for row in ws.iter_rows():
            for c in row:
                allowed=(name=='身份与说明' and c.column==2 and 6<=c.row<=14) or (name=='评分50' and 6<=c.row<=55 and 2<=c.column<=6)
                if not allowed and c.value not in (None,''):actual[c.coordinate]=c.value
        sheet_checks[name]={'shape_matches':(ws.max_row,ws.max_column)==(ref['max_row'],ref['max_column']),'fixed_cells_match':actual==ref['cells']}
        for cell in set(actual)|set(ref['cells']):
            if actual.get(cell)!=ref['cells'].get(cell):modified.append(name+'!'+cell)
    counts=Counter();missing=[];basis_required=[];ids=[]
    ws=w['评分50']
    for i in range(6,56):
        code,c,a,basis,confidence,note=[ws.cell(i,j).value for j in range(1,7)];ids.append(code)
        for name,value in [('correctness',c),('applicability',a),('basis',basis),('confidence',confidence),('note',note)]:counts[name]+=value not in (None,'')
        for j,label in [(2,'正确性'),(3,'适用性'),(5,'信心')]:
            if ws.cell(i,j).value in (None,''):missing.append({'blind_id':code,'cell':f'评分50!{ws.cell(i,j).coordinate}','field':label})
        if c in ['正确','可接受'] and (basis is None or not str(basis).strip()):basis_required.append({'blind_id':code,'cell':f'评分50!D{i}','field':'正确/可接受必须提供具体依据','original_correctness':c})
    identity=[]
    for i in range(6,15):identity.append({'cell':f'身份与说明!B{i}','field':w['身份与说明'].cell(i,1).value,'populated':w['身份与说明'].cell(i,2).value not in (None,'')})
    w.close()
    sample=read(COORD/'SAMPLE_FROZEN.json');sameids=set(ids)=={r['blind_id'] for r in sample} and len(set(ids))==50
    assert all(sha(p)==h for p,h in current.items());assert sha(raw)==before
    dump(H/'VALIDATION_REPORT.json',{'attachment':{'file_id':'file_00000000498881fd9af9aaec067d0d23','library_file_id':'libfile_9c7160f02ad48191abad4dfe51e76cfa','bytes':35725,'sha256':before},'frozen_file_hashes':checks,'sheet_checks':sheet_checks,'modified_fixed_cells':modified,'sample_ids_match':sameids,'filled_cells':dict(counts),'missing_required_ratings':missing,'missing_required_basis':basis_required,'identity_completeness':identity,'frozen_validator_exit_code':result.returncode,'formal_statistics':'blocked; not computed','reviewer_identity_and_ai_assistance':'pending author confirmation; no independent external expert claim','protected_files_unchanged_since_return_start':len(current),'raw_unchanged':True,'utc':datetime.now(timezone.utc).isoformat()})
    print(json.dumps({'filled':dict(counts),'missing_required_ratings':missing,'required_basis_missing_count':len(basis_required),'fixed_cells_modified':modified,'ids_match':sameids,'validator_issues':response['validation'][0]['issues']},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
