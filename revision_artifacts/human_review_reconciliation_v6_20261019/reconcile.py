"""Read-only human/Qwen alignment; writes only new derived outputs."""
import csv,hashlib,html,json
from collections import Counter
from fractions import Fraction
from pathlib import Path
import openpyxl
HERE=Path(__file__).resolve().parent;PROJECT=HERE.parents[1]
WORK=PROJECT/'revision_artifacts/qualitative_5arm_authorized_run_20261019'
V6=WORK/'postrun_review_v6_20261019'
OUT=HERE/'results'
read=lambda p:json.loads(p.read_text(encoding='utf8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
MAP={
 'support':{'直接支持':'direct','部分支持':'partial','相矛盾':'contradicted','证据不足':'insufficient','不确定':'uncertain'},
 'applicability':{'明确符合':'explicit_match','明确不符':'explicit_mismatch','一般规则有依据':'general_rule_with_basis','未知':'unknown'},
 'verifiability':{'可检验':'verifiable','不可检验':'not_verifiable','不确定':'uncertain'},
 'faithfulness':{'原文摘录':'quote','忠实转述':'faithful_paraphrase','新增推导或综合':'new_inference_or_synthesis','无依据':'unsupported','不确定':'uncertain'},
}
LABELS={k:list(v.values()) for k,v in MAP.items()}
ZH={'support':'支持判断','applicability':'适用条件','verifiability':'可验证性','faithfulness':'忠实性'}

def metrics(rows,dimension):
    labels=LABELS[dimension];matrix={a:{b:0 for b in labels} for a in labels};weighted={a:{b:Fraction(0) for b in labels} for a in labels}
    n=0;agree=0;den=Fraction(0);num=Fraction(0)
    for r in rows:
        h=r['human_normalized'][dimension];m=r['machine_raw'][dimension];w=Fraction(r['stratum_population'],r['stratum_sample_count'])
        matrix[m][h]+=1;weighted[m][h]+=w;n+=1;den+=w
        if m==h:agree+=1;num+=w
    return {'dimension':dimension,'matrix_orientation':'rows=Qwen original labels; columns=human normalized labels','labels':labels,'confusion_matrix':matrix,'weighted_confusion_matrix':{a:{b:float(v) for b,v in row.items()} for a,row in weighted.items()},'sample_n':n,'sample_agree':agree,'sample_agreement':agree/n,'weighted_agree_numerator':float(num),'weighted_denominator':float(den),'weighted_agreement':float(num/den),'weighted_agreement_exact':str(num/den),'formula':'sum((N_h/n_h)*I(machine==human))/sum(N_h/n_h)','interpretation':'Mechanical descriptive agreement only; scope not fully established; single reviewer, small post-run stratified sample; not engineering ground truth or validated reliability.'}

def main():
    OUT.mkdir(exist_ok=True)
    provenance=read(HERE/'ORIGINAL_PROVENANCE.json');original=Path(provenance['local_original']);before=sha(original)
    assert before=='7b831ce37bd6d8354f84ac4ef175dee7d0e8a20d43aafbf597d03fdec315ab91'
    baseline=V6/'outputs/01a13537-v6/01_盲审材料/中文盲审评分与原文.xlsx'
    old=openpyxl.load_workbook(baseline,data_only=False);new=openpyxl.load_workbook(original,data_only=False)
    assert old.sheetnames==new.sheetnames
    nonanswer_changes=[];answer_changes=[];blank_representation_changes=[]
    for sn in old.sheetnames:
        a,b=old[sn],new[sn]
        for row in range(1,max(a.max_row,b.max_row)+1):
            for col in range(1,max(a.max_column,b.max_column)+1):
                av,bv=a.cell(row,col).value,b.cell(row,col).value
                allowed=(sn=='评分30' and 6<=row<=35 and 2<=col<=8) or (sn=='异常核对' and row==6 and 2<=col<=8)
                if av!=bv and av in (None,'') and bv in (None,''):
                    blank_representation_changes.append({'sheet':sn,'cell':a.cell(row,col).coordinate,'before':av,'after':bv})
                    continue
                if av!=bv:
                    target=answer_changes if allowed else nonanswer_changes
                    target.append({'sheet':sn,'cell':a.cell(row,col).coordinate,'before':av,'after':bv})
    if nonanswer_changes:raise ValueError('Non-answer cells changed; inspect before alignment: '+json.dumps(nonanswer_changes,ensure_ascii=False))
    keys=read(V6/'outputs/01a13537-v6/02_仅协调者_勿发盲审者/本地解盲键_含机器标签.json');key={r['code']:r for r in keys}
    assert len(key)==31
    items=[];missing=[];invalid=[]
    for sn,count in [('评分30',30),('异常核对',1)]:
        for row in range(6,6+count):
            values=[new[sn].cell(row,c).value for c in range(1,9)];code=values[0];k=key[code];human={};normalized={}
            for i,dim in enumerate(MAP,1):
                human[dim]=values[i]
                if values[i] is None:missing.append({'code':code,'dimension':dim})
                elif values[i] not in MAP[dim]:invalid.append({'code':code,'dimension':dim,'value':values[i]})
                else:normalized[dim]=MAP[dim][values[i]]
            included=k['included_in_machine_agreement'];j=k.get('machine_judgment')
            if included:
                raw=read(WORK/k['response_file']);assert json.loads(raw['response']['choices'][0]['message']['content'])==j and j['review_id']==k['review_id']
            s=k.get('sampling',{})
            items.append({'code':code,'original_sheet':sn,'original_row':row,'review_id':k['review_id'],'source_id':k['source_id'],'group':k['group'],'included_in_30_agreement':included,'human_raw_zh':human,'human_normalized':normalized,'machine_raw':j,'stratum_raw_support':s.get('raw_support'),'stratum_population':s.get('population'),'stratum_sample_count':s.get('selected'),'nominal_weight':s.get('nominal_weight'),'anchor_quality':k.get('anchor_quality'),'note':values[5],'reviewer_optional':values[6],'date_optional':values[7],'scope_status':'Author reports human target is Q criterion; machine prompt does not strictly isolate target; H01 machine reason also evaluates non-Q candidate text' if code=='H01' else 'Author reports Q-only human target; machine target not precisely isolated; mechanical cross-scope description only','required_explanation_missing_in_original':code=='H16' and values[1]=='相矛盾' and not values[5]})
    assert len({r['code'] for r in items})==31 and not missing and not invalid
    scored=[r for r in items if r['included_in_30_agreement']];assert len(scored)==30
    stats={dim:metrics(scored,dim) for dim in MAP}
    assert all(v['sample_n']==30 and v['weighted_denominator']==440 for v in stats.values())
    strata=[]
    for s in read(V6/'outputs/01a13537-v6/02_仅协调者_勿发盲审者/分层总体_配额_名义概率权重.json'):
        rows=[r for r in scored if r['group']==s['group'] and r['stratum_raw_support']==s['raw_support']]
        assert len(rows)==s['selected']
        strata.append({**s,'support_agree_count':sum(r['machine_raw']['support']==r['human_normalized']['support'] for r in rows)})
    report={'original_sha256':before,'independent_parent_sha256_match':True,'library_file_id':provenance['library_file_id'],'sample_identity_and_nonanswer_cells_exact_match':not nonanswer_changes,'nonanswer_changes':nonanswer_changes,'blank_representation_changes':blank_representation_changes,'changed_answer_cells':len(answer_changes),'reviewed_items':31,'four_label_fields_completed':31,'missing_labels':missing,'invalid_labels':invalid,'duplicate_ids':0,'all_notes_blank':all(not r['note'] for r in items),'optional_reviewer_and_date_blank':True,'explicit_required_explanation_missing':['H16: F21; please identify contradictory evidence ID and exact wording required by instruction B12'],'human_support_distribution_30':dict(Counter(r['human_raw_zh']['support'] for r in scored)),'metrics':stats,'strata':strata,'exception':next(r for r in items if not r['included_in_30_agreement']),'scope_note':'Both workflows show the complete candidate. Frozen Qwen prompt evaluates one anonymous candidate and does not restrict support to Q alone. Human instructions say all substantive criteria but do not fully specify Q-only vs whole-candidate evaluation. Mechanical alignment is possible; same-scope interpretation remains conditional, especially H01.','sequence_note':'User stated review was already filled. Actual filling sequence relative to later examples is not documented; no contamination finding is inferred from later questions alone.','limitations':['Single reviewer; no inter-rater reliability estimate.','30 small post-run stratified items; weights are nominal N_h/n_h under fixed-hash sampling interpretation.','No merged direct+partial accuracy; agreement is not factual correctness.','H16 label retained despite missing required conflict explanation.','X01 diagnosis remains separate; no Qwen label manufactured.'],'new_model_api_calls':0}
    supplement=read(OUT/'作者补充说明.json')
    report['author_supplement']=supplement
    report['original_required_explanation_missing']=report.pop('explicit_required_explanation_missing')
    report['remaining_author_questions']=[]
    report['scope_note']=supplement['interpretation']
    report['limitations']=[x for x in report['limitations'] if not x.startswith('H16 label')] + ['H16 original note remains blank; author supplied separate explanation at2026-10-19T08:55:40Z, without changing any original score.']
    for r in items:r['author_supplement_reference']='作者补充说明.json';r['requires_further_author_answer']=False
    for name,value in [('REPORT.json',report),('ITEM_ALIGNMENT.json',items),('ORIGINAL_ANSWER_CELLS.json',answer_changes)]:
        (OUT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2,default=str),encoding='utf8')
    with (OUT/'逐项人机对照.csv').open('w',encoding='utf-8-sig',newline='') as f:
        fields=['编号','计入30条','组别','源需求','Qwen支持原标签','人工支持原值','支持是否一致','Qwen适用性','人工适用性','Qwen可验证性','人工可验证性','Qwen忠实性','人工忠实性','层总数','层样本数','名义权重','锚点质量','原备注','口径或说明事项']
        writer=csv.writer(f);writer.writerow(fields)
        for r in items:
            j=r['machine_raw'] or {};h=r['human_raw_zh'];note='作者说明人工只审[Q]；机器H01理由涉及候选其他字段' if r['code']=='H01' else '原F21空白；作者已于08:55:40UTC另补理由，未改原件' if r['required_explanation_missing_in_original'] else 'X01独立缺失诊断，不计入一致率' if not r['included_in_30_agreement'] else ''
            writer.writerow([r['code'],r['included_in_30_agreement'],r['group'],r['source_id'],j.get('support'),h['support'],j.get('support')==r['human_normalized']['support'] if j else '',j.get('applicability'),h['applicability'],j.get('verifiability'),h['verifiability'],j.get('faithfulness'),h['faithfulness'],r['stratum_population'],r['stratum_sample_count'],r['nominal_weight'],(r['anchor_quality'] or {}).get('status'),r['note'],note])
    h01=next(r for r in items if r['code']=='H01');ex=report['exception']
    lines=['v6人工审核核验与机械人机对照','', '真实完成范围：H01–H30及X01，共31项，四个标签栏全部已填；无缺填标签、无非法标签、无重复编号。备注全空，姓名/日期为可选项，不能据此判整份未完成。','文件一致性：样本编号、原需求、候选、判据、实际引用原文、规则和评分字段与最终v6一致。124个标签答案单元格改变；另有空单元格None/空字符串的存储表示变化，属于空白分隔行，不涉及证据或答案变更。原件保持不变。','SHA256：'+before,'','30条机械一致率（非可靠性验证或领域真值）：']
    for dim,s in stats.items():lines.append(f"{ZH[dim]}：未加权 {s['sample_agree']}/30 = {s['sample_agreement']:.2%}；名义分层加权 {s['weighted_agree_numerator']:.1f}/440 = {s['weighted_agreement']:.2%}。")
    lines+=['权重为原机器支持标签×生成臂层的N_h/5；各维度复用同一抽样权重，不按人工答案重新分层。加权仅描述性，不恢复随机化或消除事后抽样偏差。','人工支持分布：'+json.dumps(report['human_support_distribution_30'],ensure_ascii=False),'','需要作者确认的两点：','1. H01评价单位是只评价[Q]判据，还是包含整个候选/验收要求？双方均展示完整候选，但具体评价单位未完全明示。H01机器原值='+h01['machine_raw']['support']+'，人工原值='+h01['human_raw_zh']['support']+'。双方支持标签同为部分支持，适用性和可验证性也一致；忠实性为机器“新增推导或综合”、人工“忠实转述”。机器理由明确提到整体候选中的温度控制要求/配置协调。人审备注为空，不能据此确定人审范围。不先判定哪方错误，不将忠实性差异直接解释为同口径分歧。','2. H16维持“相矛盾”。模板B12明确要求注明冲突证据编号和具体文字，而F21为空。请补充这一必需解释；不从孤立摘录自行推断矛盾。','其余备注为空属于解释能力限制；不要求整份重做，不把可选姓名和日期当缺失必填。','','X01单独诊断：'+json.dumps(ex['human_raw_zh'],ensure_ascii=False)+'。原机器ID异常仍无效，人工判断不替代机器缺失，X01不进入任何30条一致率。','','限制：单人审核、30条小样本、结果揭晓后的分层抽样，不能称为领域事实真值，也不估计评审者间一致性。direct与partial不合并为正确率。','填写顺序：用户已明确表示此前填过；实际填写与后续询问例子的先后细节未独立记录，不能仅凭后来询问例子宣布污染。','原始覆盖仍为441候选、440机器有效、1机器异常；本次无新增模型调用，无答案改写，无论文整合，无重开S2。']
    lines=[x for x in lines if not x.startswith(('需要作者确认的两点：','1. H01','2. H16'))]
    lines += ['','作者补充（2026-10-19 08:55:40 UTC，单独记录，不回填原件）：','人工目标已由作者明确为[Q]那句话；H16相矛盾理由已补充，见作者补充说明.json。不再要求两项确认，也不要求重做人审。','H16原工作簿F21仍为空，此补充仅归于作者转述，不伪造专家签名或原填写记录。','H01支持双方均为部分支持，适用性和可验证性一致；忠实性为机器新增推导/综合、人工忠实转述。机器理由涉及[Q]外候选内容，故不能把忠实性差异直接认定为同口径误判。','机器输入包含完整候选不等于已严格规定对整条逐字段评分。冻结提示词未将目标明确限定只审[Q]；结合H01理由，同口径无法证明。因此四项一致率均只作跨口径机械描述，不能称准确率或机器评审可靠性通过人工验证。','目前无待作者回答事项。']
    (OUT/'中文简明汇总.txt').write_text('\n'.join(lines)+'\n',encoding='utf8')
    esc=html.escape;cards=[]
    for dim,s in stats.items():
        table='<table><tr><th>机器原标签 / 人工归一标签</th>'+''.join('<th>'+esc(l)+'</th>' for l in s['labels'])+'</tr>'
        for label,row in s['confusion_matrix'].items():table+='<tr><th>'+label+'</th>'+''.join('<td>'+str(row[x])+'</td>' for x in s['labels'])+'</tr>'
        cards.append('<h2>'+ZH[dim]+'混淆矩阵（计数）</h2>'+table+'</table>')
    with (OUT/'逐项人机对照.csv').open(encoding='utf-8-sig',newline='') as f:table=list(csv.reader(f))
    grid='<table><thead><tr>'+''.join('<th>'+esc(x)+'</th>' for x in table[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(x)+'</td>' for x in row)+'</tr>' for row in table[1:])+'</tbody></table>'
    document='<!doctype html><meta charset="utf-8"><title>人工审核机械对照</title><style>body{font:16px/1.65 "Microsoft YaHei",sans-serif;padding:28px;color:#21334b}pre{white-space:pre-wrap}table{border-collapse:collapse;margin:18px 0;font-size:14px}th,td{border:1px solid #ccd4df;padding:8px;vertical-align:top}th{background:#edf2f8}section{overflow:auto}</style><h1>人工审核核验与机械对照</h1><pre>'+esc('\n'.join(lines))+'</pre>'+''.join(cards)+'<h2>逐项对照（含解盲信息）</h2><section>'+grid+'</section>'
    (OUT/'中文汇总与逐项对照.html').write_text(document,encoding='utf8')
    assert sha(original)==before
    print(json.dumps({'matched':True,'labels_completed':31,'metrics':{k:{kk:v[kk] for kk in ['sample_agree','sample_agreement','weighted_agree_numerator','weighted_agreement']} for k,v in stats.items()},'H01_machine':h01['machine_raw'],'H01_human':h01['human_raw_zh'],'required_note_missing_in_original':['H16'],'author_supplement_received':True,'remaining_author_questions':[],'new_calls':0},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
