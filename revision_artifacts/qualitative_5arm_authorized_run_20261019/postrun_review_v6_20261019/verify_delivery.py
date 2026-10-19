import json,re,html,hashlib,zipfile
from pathlib import Path
from collections import Counter
from html.parser import HTMLParser
import openpyxl
from sample_v6 import select
HERE=Path(__file__).resolve().parent;ROOT=HERE/'outputs/01a13537-v6';BLIND=ROOT/'01_盲审材料';PRIVATE=ROOT/'02_仅协调者_勿发盲审者'
read=lambda p:json.loads(p.read_text(encoding='utf8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
class Reader(HTMLParser):
    def __init__(self):super().__init__();self.quotes=[];self.current=None;self.external=[];self.script=False;self.ids=[]
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if tag=='pre' and a.get('class')=='quote':self.current=''
        if tag=='article':self.ids.append(a['id'])
        if tag=='script':self.script=True
        for k in ['href','src']:
            if a.get(k,'').startswith(('http:','https:','//')):self.external.append(a[k])
    def handle_data(self,data):
        if self.current is not None:self.current+=data
    def handle_endtag(self,tag):
        if tag=='pre' and self.current is not None:self.quotes.append(self.current);self.current=None

def main():
    items=read(HERE/'blind_input.json');expected=read(HERE/'reader_rows.json')['rows']
    xlsx=BLIND/'中文盲审评分与原文.xlsx';wb=openpyxl.load_workbook(xlsx,data_only=False)
    assert wb.sheetnames==['使用说明','评分30','异常核对','原文阅读']
    assert all(s.sheet_state=='visible' for s in wb)
    for sheet,count in [('评分30',30),('异常核对',1)]:
        s=wb[sheet];assert all(s.cell(row,col).value is None for row in range(6,6+count) for col in range(2,9))
        assert len(s.data_validations.dataValidation)==4
        assert s.freeze_panes=='A6'
    assert [wb['评分30'].cell(i+6,1).value for i in range(30)]==[b['code'] for b in items[:30]]
    assert wb['异常核对']['A6'].value=='X01'
    for i,row in enumerate(expected,6):
        for col,value in enumerate(row,1):assert (wb['原文阅读'].cell(i,col).value or '')==value,(i,col)
    # Reassemble Excel's consecutive excerpt chunks and match the exact frozen packets.
    reconstructed={}
    for code,label,text in expected:
        if label.startswith('证据原文 '):
            eid=label.split(' ')[1].split('（')[0];reconstructed.setdefault((code,eid),'');reconstructed[code,eid]+=text
    quotes=[]
    for b in items:
        for e in b['cited_evidence']:
            assert reconstructed[b['code'],e['evidence_id']]==e['supplied_excerpt'];quotes.append(e['supplied_excerpt'])
    page=(BLIND/'离线阅读.html').read_text(encoding='utf8');parser=Reader();parser.feed(page)
    assert parser.quotes==quotes and parser.ids==[b['code'] for b in items]
    assert not parser.script and not parser.external
    with zipfile.ZipFile(xlsx) as archive:xml='\n'.join(archive.read(n).decode('utf8') for n in archive.namelist() if n.endswith('.xml'))
    for content in [xml,page,(BLIND/'阅读与保存说明.txt').read_text(encoding='utf8')]:
        assert not re.search(r'semantic_new|semantic_old|random_new|random_old|norag_original|deepseek|qwen|OR\d{3}|[0-9a-f]{64}',content,re.I),'explicit_blinding_leak'
    policy=read(PRIVATE/'协议修订_v6.json');keys=read(PRIVATE/'本地解盲键_含机器标签.json');strata=read(PRIVATE/'分层总体_配额_名义概率权重.json')
    allrows=read(HERE.parent/'authorized_remaining330_20261019/terminal_results/REVIEW_DIAGNOSTICS_441.json');valid=[r for r in allrows if r['review_status']=='completed']
    selected,again=select(valid,[r['review_id'] for r in valid],policy)
    assert strata==again and {r['review_id'] for r in selected}=={r['review_id'] for r in keys if r['included_in_machine_agreement']}
    assert len(keys)==31 and sum(r['included_in_machine_agreement'] for r in keys)==30
    assert sum(s['population'] for s in strata)==440 and sum(s['selected'] for s in strata)==30
    for s in strata:
        if s['population']:assert abs(s['nominal_inclusion_fraction']*s['nominal_weight']-1)<1e-12
        else:assert s['selected']==0 and s['nominal_inclusion_fraction'] is None
    result={'spreadsheet_sheets':4,'blank_scoring_records':31,'sample_size':30,'exception_count':1,'nominal_population':440,'all_candidate_coverage':441,'evidence_excerpts_verified_exact_excel_and_html':len(quotes),'reader_rows_verified':len(expected),'explicit_group_model_OR_internal_id_leaks':0,'hidden_sheets':0,'html_external_resources':0,'html_script_or_fake_save':False,'raw_normalized_mapping':'identity; Chinese display mapping only','reproducible_sample':True,'native_excel_interactive_save_test':'not performed; standard XLSX export read back with native list validations and blank input cells','api_calls':0}
    (HERE/'DELIVERY_VALIDATION.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':main()
