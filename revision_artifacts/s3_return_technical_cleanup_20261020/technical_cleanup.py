"""Lossless ZIP-part normalization and strict dimension-compatible frozen validation."""
import copy,hashlib,importlib.util,json,re,zipfile
from pathlib import Path
from unittest.mock import patch
import openpyxl
H=Path(__file__).resolve().parent;ROOT=H.parents[1];F=ROOT/'revision_artifacts/s3_paired_review50_20261019';C=F/'02_协调者_勿发评审'
RAW=ROOT/'revision_artifacts/s3_return_analysis_20261020_v2/raw/评分表.xlsx'
OLD=ROOT/'revision_artifacts/s3_return_analysis_20261020/raw/数值候选独立评分.xlsx'
TEMPLATE=F/'outputs/01a13537-s3/01_仅此目录可发评审/数值候选独立评分.xlsx'
OUT=H/'normalized/评分表_仅分隔格技术整理.xlsx'
read=lambda p:json.loads(p.read_text(encoding='utf8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf8')
spec=importlib.util.spec_from_file_location('frozen_s3_validation',F/'analyze.py');frozen=importlib.util.module_from_spec(spec);spec.loader.exec_module(frozen)
ORIGINAL_LOADER=openpyxl.load_workbook

def measured_loader(*args,**kwargs):
    wb=ORIGINAL_LOADER(*args,**kwargs)
    if kwargs.get('read_only'):
        for ws in wb:
            ws.reset_dimensions();ws.calculate_dimension(force=True)
    return wb

def compatible_validate(path):
    freeze=read(C/'ANALYSIS_FREEZE.json')
    assert all(sha(F/name)==h for name,h in freeze['files'].items())
    expected=copy.deepcopy(read(C/'WORKBOOK_CONTRACT.json'))
    wb=measured_loader(TEMPLATE,read_only=True,data_only=False)
    for ws in wb:
        for key,value in [('max_row',ws.max_row),('max_column',ws.max_column)]:
            if expected[ws.title][key] is None:expected[ws.title][key]=value
            else:assert expected[ws.title][key]==value
    wb.close()
    # Only derive actual extents; the exact frozen content, formula, missing-data and declaration checks run unchanged.
    with patch.object(frozen.openpyxl,'load_workbook',measured_loader):
        return frozen.validate(path,expected,read(C/'SAMPLE_FROZEN.json'))

def snapshot(path):
    wb=ORIGINAL_LOADER(path,read_only=True,data_only=False);result={}
    for ws in wb:
        result[ws.title]={c.coordinate:(c.value,c.data_type) for row in ws.iter_rows() for c in row if c.value not in (None,'')}
    wb.close();return result

def main():
    assert sha(RAW)=='f5f14e175f255495c2187b00311a0348bc4a977e8b5532308cf10f3fa9b5773b'
    assert sha(OLD)=='87f3153b7ddb329efd528a45621135a812298c0ca812d42eb59caae6a03a173f'
    changes=[x for x in read(ROOT/'revision_artifacts/s3_return_analysis_20261020_v2/CELL_DIFF.json') if x['sheet']=='完整原文']
    assert len(changes)==150 and all(x['old'] in ('',None) and str(x['new'])=='136' for x in changes)
    before=snapshot(RAW);old=snapshot(OLD);template=snapshot(TEMPLATE)
    for x in changes:
        cell=x['cell'];assert cell not in old['完整原文'] and cell not in template['完整原文'];assert str(before['完整原文'][cell][0])=='136'
    assert not OUT.exists();OUT.parent.mkdir(exist_ok=True)
    with zipfile.ZipFile(RAW) as zin,zipfile.ZipFile(OUT,'w') as zout:
        for info in zin.infolist():
            payload=zin.read(info.filename)
            if info.filename=='xl/worksheets/sheet4.xml':
                text=payload.decode('utf8')
                for x in changes:
                    cell=x['cell'];pattern=r'(<(?:\w+:)?c\b[^>]*\br="'+re.escape(cell)+r'"[^>]*>)(.*?)(</(?:\w+:)?c>)'
                    matches=list(re.finditer(pattern,text,re.S));assert len(matches)==1,cell
                    m=matches[0];opening=re.sub(r'\s+t="[^"]*"','',m[1])
                    # Preserve cell style/attributes and all other XML bytes; only remove its value representation.
                    text=text[:m.start()]+opening+m[3]+text[m.end():]
                payload=text.encode('utf8')
            zout.writestr(info,payload)
    after=snapshot(OUT);expected=copy.deepcopy(before)
    for x in changes:del expected['完整原文'][x['cell']]
    assert after==expected
    with zipfile.ZipFile(RAW) as a,zipfile.ZipFile(OUT) as b:
        changed_parts=[n for n in a.namelist() if a.read(n)!=b.read(n)]
        assert a.namelist()==b.namelist() and changed_parts==['xl/worksheets/sheet4.xml']
    dump(H/'NORMALIZATION_LOG.json',{'input_sha256':sha(RAW),'output_sha256':sha(OUT),'changed_zip_parts':changed_parts,'cells':[dict(sheet=x['sheet'],cell=x['cell'],before=x['new'],after=None,proof='Blank in frozen template and v1 attachment') for x in changes],'all_other_nonempty_cell_values_and_types_unchanged':True,'rating_basis_note_identity_cells_unchanged':True})
    expected_contract=read(C/'WORKBOOK_CONTRACT.json');sample=read(C/'SAMPLE_FROZEN.json')
    for label,path in [('raw_v2',RAW),('normalized',OUT)]:
        original=frozen.validate(path,expected_contract,sample);compat=compatible_validate(path)
        dump(H/(label+'_validation.json'),{'original_frozen':{k:original[k] for k in ['valid','issues','input_sha256']},'dimension_compatible':{k:compat[k] for k in ['valid','issues','input_sha256']}})
    assert not compatible_validate(OUT)['valid']
    print('150 separator values cleared only in separate copy; real declaration gaps remain blocked.')
if __name__=='__main__':main()
